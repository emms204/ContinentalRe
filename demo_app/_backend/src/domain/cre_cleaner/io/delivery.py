"""Delivery of a multi-file batch (IMPL-20260929-06, reworked IMPL-20261002-04).

Grouping always comes first (:mod:`cre_cleaner.core.batch`): every delivery
option hands over the per-group cleaned workbooks exactly as the pipeline
wrote them — one upload-ready workbook per cedant / broker / year / quarter
(and per currency where the pipeline splits), with the normal period title,
the SUMMARY sheet and the template S/NO. Nothing is rewritten here.

* ``separate`` (default): those workbooks, one file each.
* ``zip``: exactly those workbooks in one archive.
* ``merged``: one workbook when the batch produced exactly one; otherwise the
  same archive as ``zip``. A multi-period workbook is never built.

Provenance (groups, source files, rows per sheet, columns that differ between
groups) is returned by :func:`delivery_report` for the app / API response and
logs — it is never written into a deliverable.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

DELIVERY_OPTIONS = ("separate", "zip", "merged")
_SKIP_SHEETS = {"SUMMARY", "EXCEPTIONS", "SOURCE AUDIT", "SOURCE_AUDIT"}
_SERIAL_HEADERS = {"S/NO", "S/N", "SNO", "SN"}
# Period lists longer than this are named first_to_last in the zip name.
_MAX_NAMED_PERIODS = 6


def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().upper().rstrip(".")


def _safe(v: Any) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", str(v or "")).strip("_").upper()


def header_row_index(ws, max_scan: int = 12) -> Optional[int]:
    """1-based header row of a cleaned sheet (the row with the S/NO cell)."""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(r, c).value) in _SERIAL_HEADERS:
                return r
    return None


def data_rows(ws) -> List[int]:
    """Row numbers of data rows (below the header, any non-empty cell)."""
    h = header_row_index(ws)
    if h is None:
        return []
    out = []
    for r in range(h + 1, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            out.append(r)
    return out


def _sheet_headers(path: Path) -> Dict[str, Tuple[int, List[str]]]:
    """Data sheet title → (data row count, header names) of one workbook."""
    wb = load_workbook(path)
    try:
        out: Dict[str, Tuple[int, List[str]]] = {}
        for ws in wb.worksheets:
            if ws.title.upper() in _SKIP_SHEETS:
                continue
            h = header_row_index(ws)
            if h is None:
                continue
            hdr = [_norm(ws.cell(h, c).value) for c in range(1, ws.max_column + 1)]
            out[ws.title] = (len(data_rows(ws)), [v for v in hdr if v])
        return out
    finally:
        wb.close()


def data_sheet_counts(path: Path) -> Dict[str, int]:
    return {k: n for k, (n, _h) in _sheet_headers(Path(path)).items()}


def upload_files(outputs: Sequence[Dict[str, Any]]) -> List[Path]:
    """The per-group upload-ready workbooks (existing files, in order, unique)."""
    seen, out = set(), []
    for o in outputs:
        p = Path(str(o.get("output_path") or ""))
        if p.is_file() and str(p) not in seen:
            seen.add(str(p))
            out.append(p)
    return out


def _periods(outputs: Sequence[Dict[str, Any]]) -> List[Tuple[int, int]]:
    out = set()
    for o in outputs:
        try:
            out.add((int(o.get("year")), int(o.get("quarter"))))
        except (TypeError, ValueError):
            continue
    return sorted(out)


def zip_name(outputs: Sequence[Dict[str, Any]], *, stem: Optional[str] = None) -> str:
    """Clear archive name: ``<CEDANT>_<BROKER>_<YEAR>_Q#_Q#_cleaned.zip``.

    Periods of one year share the year (``2024_Q2_Q4``); several years list
    each period (``2021_Q4_2022_Q1``) or, past a few, first ``to`` last.
    Falls back to ``stem`` when the outputs carry no cedant / period.
    """
    parties = []
    for o in outputs:
        pair = (_safe(o.get("cedant")), _safe(o.get("broker")))
        if pair[0] and pair not in parties:
            parties.append(pair)
    periods = _periods(outputs)
    if not parties or not periods:
        return f"{_safe(stem) or 'BATCH'}_cleaned_workbooks.zip"
    who = "_".join(f"{c}_{b}" if b else c for c, b in parties) if len(parties) <= 2 else "MULTI_PARTNER"
    years = sorted({y for y, _q in periods})
    if len(years) == 1:
        when = f"{years[0]}_" + "_".join(f"Q{q}" for _y, q in periods)
    elif len(periods) <= _MAX_NAMED_PERIODS:
        when = "_".join(f"{y}_Q{q}" for y, q in periods)
    else:
        (y0, q0), (y1, q1) = periods[0], periods[-1]
        when = f"{y0}_Q{q0}_to_{y1}_Q{q1}"
    return f"{who}_{when}_cleaned.zip"


def build_zip(outputs: Sequence[Dict[str, Any]], dest: Path) -> Path:
    """Zip exactly the per-group cleaned workbooks (flat, by file name)."""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    files = upload_files(outputs)
    names = [p.name for p in files]
    if len(set(names)) != len(names):
        raise ValueError(f"duplicate workbook names in batch: {names}")
    with zipfile.ZipFile(dest, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for p in files:
            zf.write(p, arcname=p.name)
    return dest


def deliver(
    outputs: Sequence[Dict[str, Any]],
    option: str,
    dest_dir: Path,
    *,
    stem: Optional[str] = None,
) -> List[Path]:
    """Materialise a delivery option; returns the file(s) to hand over.

    ``separate``: every per-group workbook. ``zip``: one archive of them.
    ``merged``: the single workbook when there is exactly one, else the zip.
    """
    option = (option or "separate").strip().lower()
    if option not in DELIVERY_OPTIONS:
        raise ValueError(f"delivery option must be one of {DELIVERY_OPTIONS}, got {option!r}")
    files = upload_files(outputs)
    if option == "separate" or (option == "merged" and len(files) <= 1):
        return files
    return [build_zip(outputs, Path(dest_dir) / zip_name(outputs, stem=stem))]


def delivery_report(
    outputs: Sequence[Dict[str, Any]],
    option: str,
    handed: Sequence[Path],
    *,
    read_sheets: bool = True,
) -> Dict[str, Any]:
    """Provenance of a delivery for the app / API response / logs.

    Per workbook: group, period, currency, source files, row counts and —
    with ``read_sheets`` (opens each workbook) — rows per data sheet and
    ``columns_differing``: header names on a sheet that the same-named sheet
    of the first workbook does not have (layout drift between groups).
    """
    first_headers: Dict[str, List[str]] = {}
    workbooks: List[Dict[str, Any]] = []
    for o in outputs:
        p = Path(str(o.get("output_path") or ""))
        if not p.is_file():
            continue
        sheets = _sheet_headers(p) if read_sheets else {}
        differing: Dict[str, List[str]] = {}
        for name, (_n, hdr) in sheets.items():
            base = first_headers.setdefault(name, hdr)
            extra = [h for h in hdr if h not in base]
            if extra:
                differing[name] = extra
        workbooks.append({
            "workbook": p.name, "group": o.get("group") or "",
            "cedant": o.get("cedant"), "broker": o.get("broker"),
            "year": o.get("year"), "quarter": o.get("quarter"),
            "currency": o.get("currency") or "",
            "source_files": list(o.get("source_files") or []),
            "sheet_rows": {k: n for k, (n, _h) in sheets.items()},
            "premium_rows": int(o.get("premium_rows") or 0),
            "claims_rows": int(o.get("claims_rows") or 0),
            "outstanding_rows": int(o.get("outstanding_rows") or 0),
            "columns_differing": differing,
        })
    handed = [Path(h) for h in handed]
    return {
        "delivery": (option or "separate").strip().lower(),
        "returned": [h.name for h in handed],
        "kind": "zip" if any(h.suffix.lower() == ".zip" for h in handed) else "workbook",
        "periods": [f"{y} Q{q}" for y, q in _periods(outputs)],
        "workbooks": workbooks,
    }
