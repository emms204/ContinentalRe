"""Upload-structure validator for cleaned workbooks (IMPL-20261006-03 B5).

Checks one written workbook against the upload contract in ``config``:

  (a) SUMMARY is the first sheet; every other upload sheet is
      ``<Class> - PREMIUM|CLAIMS|OUTSTANDING`` (or the legacy collapsed names),
      and no sheet is a ``Facultative`` class sheet;
  (b) the header row is exactly the contract headers, in order
      (premium row 3 from column A, claims / outstanding row 4 from column B
      with column A empty);
  (c) no duplicate or blank header;
  (d) nothing right of the last header column (or in claims column A);
  (e) nothing hidden (sheets, rows, columns);
  (f) no whitespace-only text cell;
  (g) no Excel error literal (``#N/A``, ``#REF!`` …);
  (h) the band row carries the contract band labels.

``template_header_deviations`` compares the reference ``TEMPLATE.xlsx`` with
the contract after aligning it (IMPL-20261006-04): the premium CHANNEL / SUB
CHANNEL columns are dropped (Cleaning Manual Part Four Step 12 and Part Seven:
never carried into the cleaned Premium output) and text is compared with
whitespace normalised. What remains is listed in ``TEMPLATE_DEVIATIONS_ALLOWED``
(the CHANNEL exclusion and whitespace-only items); anything else is returned
by ``template_alignment_problems``. ``validate_against_template`` checks a
written workbook's header rows and band labels against the aligned template.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from src.domain.cre_cleaner.config import (
    CLAIMS_COL_MAP,
    CLAIMS_HEADERS,
    DEFAULT_PROPORTION_MODE,
    FAC_CLASS_LABEL,
    PREMIUM_BAND_LABELS,
    TEMPLATE_SHEETS,
    premium_headers,
)

PREMIUM_HEADER_ROW = 3
CLAIMS_HEADER_ROW = 4
_TYPES = ("PREMIUM", "CLAIMS", "OUTSTANDING")
_AUDIT_SHEETS = frozenset({"SOURCE AUDIT", "EXCEPTIONS"})
_EXCEL_ERRORS = frozenset({
    "#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!",
    "#GETTING_DATA", "#SPILL!", "#CALC!",
})

# Premium TEMPLATE columns that are never part of the cleaned output (Cleaning
# Manual Part Four Step 12 "Remove Channel and Sub Channel" and Part Seven
# "Channel and Sub Channel are not part of the cleaned output").
TEMPLATE_EXCLUDED_HEADERS = frozenset({"CHANNEL", "SUB CHANNEL"})

# Where the reference TEMPLATE (Template.xlsx, IMPL-20261006-04) still differs
# from the upload contract once aligned. (sheet kind, what) → why it is kept.
# Only the Manual's CHANNEL exclusion and whitespace-only items remain; header
# text and column positions of the output are otherwise identical.
TEMPLATE_DEVIATIONS_ALLOWED: Dict[Tuple[str, str], str] = {
    ("premium", "CHANNEL/SUB CHANNEL columns D:E"):
        "excluded from the cleaned Premium output (Manual Part Four Step 12, Part Seven); "
        "the TEMPLATE columns right of them sit two places further right",
    ("premium", "whitespace-only cells S1, B2, C2"):
        "validator forbids whitespace-only cells",
    ("premium", "band label leading spaces ('    R E T E N T I O N', '   F A C U L T A T I V E')"):
        "whitespace only; the contract writes the band labels without leading spaces",
    ("claims", "'DATE OF LOSS  (day-mth-year)' double space"):
        "whitespace only; contract / gold use one space",
    ("claims", "'DETAILS OF LOSS ' trailing space"):
        "whitespace only; contract / gold have no trailing space",
}


def _norm(v: Any) -> str:
    return "" if v is None else str(v)


def _sheet_kind(title: str, collapsed_names: Dict[str, str]) -> Tuple[Optional[str], str]:
    """(kind, class label) for an upload sheet title; kind None if unknown."""
    if title in collapsed_names:
        return collapsed_names[title], ""
    for t in _TYPES:
        suffix = f" - {t}"
        if title.endswith(suffix) and len(title) > len(suffix):
            kind = "premium" if t == "PREMIUM" else ("claims" if t == "CLAIMS" else "outstanding")
            return kind, title[: -len(suffix)]
    return None, ""


def _expected_headers(kind: str, proportion_mode: str, leading_blank: bool) -> Tuple[int, int, List[str]]:
    """(header row, first column, headers)."""
    if kind == "premium":
        return PREMIUM_HEADER_ROW, 1, list(premium_headers(proportion_mode))
    first = min(CLAIMS_COL_MAP.values()) if leading_blank else 1
    return CLAIMS_HEADER_ROW, first, list(CLAIMS_HEADERS)


def validate_upload_workbook(
    wb: Any,
    *,
    proportion_mode: str = DEFAULT_PROPORTION_MODE,
    claims_leading_blank: bool = True,
    allow_audit_sheets: bool = False,
) -> List[str]:
    """Problems found in an openpyxl workbook (empty list = valid)."""
    problems: List[str] = []
    names = list(wb.sheetnames)
    if not names or names[0] != "SUMMARY":
        problems.append(f"(a) first sheet is {names[0] if names else None!r}, expected 'SUMMARY'")
    collapsed_names = {TEMPLATE_SHEETS[k][0]: k for k in ("premium", "claims", "outstanding")}
    for ws in wb.worksheets:
        title = ws.title
        if getattr(ws, "sheet_state", "visible") != "visible":
            problems.append(f"(e) sheet {title!r} is {ws.sheet_state}")
        if title == "SUMMARY" or (allow_audit_sheets and title in _AUDIT_SHEETS):
            problems.extend(_cell_problems(ws, title, 1, ws.max_column or 1))
            continue
        kind, label = _sheet_kind(title, collapsed_names)
        if kind is None:
            problems.append(f"(a) unexpected sheet {title!r}")
            continue
        if label.strip().upper() == FAC_CLASS_LABEL.upper():
            problems.append(f"(a) {title!r}: Facultative is never its own class sheet")
        hrow, first, expected = _expected_headers(kind, proportion_mode, claims_leading_blank)
        last = first + len(expected) - 1
        got = [_norm(ws.cell(hrow, first + i).value) for i in range(len(expected))]
        if got != expected:
            diff = [
                f"{ws.cell(hrow, first + i).coordinate}={g!r} (want {e!r})"
                for i, (g, e) in enumerate(zip(got, expected)) if g != e
            ]
            problems.append(f"(b) {title!r} header row {hrow}: " + "; ".join(diff[:6]))
        seen: Dict[str, int] = {}
        for g in got:
            if not g.strip():
                problems.append(f"(c) {title!r}: blank header")
            seen[g] = seen.get(g, 0) + 1
        dups = sorted(k for k, n in seen.items() if n > 1 and k.strip())
        if dups:
            problems.append(f"(c) {title!r}: duplicate headers {dups}")
        if kind == "premium":
            for col, text in PREMIUM_BAND_LABELS.items():
                if _norm(ws.cell(hrow - 1, col).value) != text:
                    problems.append(f"(h) {title!r}: band {ws.cell(hrow - 1, col).coordinate} "
                                    f"is {ws.cell(hrow - 1, col).value!r}, want {text!r}")
        else:
            band_col = first + CLAIMS_HEADERS.index("FROM")
            if _norm(ws.cell(hrow - 1, band_col).value) != "INSURANCE PERIOD":
                problems.append(f"(h) {title!r}: period band missing at "
                                f"{ws.cell(hrow - 1, band_col).coordinate}")
        problems.extend(_cell_problems(ws, title, first, last))
    return problems


def _cell_problems(ws: Any, title: str, first: int, last: int) -> List[str]:
    out: List[str] = []
    outside: List[str] = []
    for (r, c), cell in list(getattr(ws, "_cells", {}).items()):
        v = cell.value
        if v is None:
            continue
        if c > last or c < first:
            outside.append(cell.coordinate)
            continue
        if isinstance(v, str):
            if not v.strip():
                out.append(f"(f) {title}!{cell.coordinate} is whitespace only")
            elif v.strip().upper() in _EXCEL_ERRORS:
                out.append(f"(g) {title}!{cell.coordinate} = {v.strip()!r}")
    if outside:
        out.append(f"(d) {title!r}: {len(outside)} cell(s) outside columns "
                   f"{first}..{last}, e.g. {sorted(outside)[:3]}")
    for key, dim in list(ws.column_dimensions.items()):
        if dim.hidden:
            out.append(f"(e) {title!r}: column {key} hidden")
    for key, dim in list(ws.row_dimensions.items()):
        if dim.hidden:
            out.append(f"(e) {title!r}: row {key} hidden")
    return out


def validate_upload_file(path: Path, **kwargs: Any) -> List[str]:
    from openpyxl import load_workbook
    wb = load_workbook(Path(path))
    try:
        return validate_upload_workbook(wb, **kwargs)
    finally:
        wb.close()


def _ws(v: Any) -> str:
    """Text with whitespace normalised (strip + single spaces)."""
    return " ".join(_norm(v).split())


def _aligned_template(ws: Any, kind: str) -> Tuple[List[Tuple[int, str]], List[int], Dict[int, int]]:
    """For one TEMPLATE sheet: ([(template col, header text)] kept after the
    exclusion, [excluded template cols], {template col: contract col})."""
    _sheet, hrow = TEMPLATE_SHEETS[kind]
    _row, first, _expected = _expected_headers(kind, DEFAULT_PROPORTION_MODE, True)
    row = [_norm(c.value) for c in ws[hrow]]
    while row and not row[-1].strip():
        row.pop()
    kept: List[Tuple[int, str]] = []
    excluded: List[int] = []
    colmap: Dict[int, int] = {}
    for col in range(first, len(row) + 1):
        text = row[col - 1]
        if kind == "premium" and _ws(text).upper() in TEMPLATE_EXCLUDED_HEADERS:
            excluded.append(col)
            continue
        colmap[col] = first + len(kept)
        kept.append((col, text))
    return kept, excluded, colmap


def template_header_deviations(template_path: Path) -> Dict[str, List[Tuple[str, str, str]]]:
    """Header differences between a TEMPLATE workbook and the contract once
    aligned: {kind: [(template cell, template text, contract text)]}. Excluded
    CHANNEL / SUB CHANNEL columns are listed with contract text '(excluded)';
    every other entry is a text difference at the aligned position."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(Path(template_path), read_only=False)
    out: Dict[str, List[Tuple[str, str, str]]] = {}
    try:
        for kind in ("premium", "claims", "outstanding"):
            sheet, hrow = TEMPLATE_SHEETS[kind]
            if sheet not in wb.sheetnames:
                out[kind] = [("-", "sheet missing", sheet)]
                continue
            kept, excluded, _colmap = _aligned_template(wb[sheet], kind)
            _row, first, expected = _expected_headers(kind, DEFAULT_PROPORTION_MODE, True)
            diffs = [(f"{get_column_letter(c)}{hrow}", _norm(wb[sheet].cell(hrow, c).value),
                      "(excluded)") for c in excluded]
            for i in range(max(len(kept), len(expected))):
                if i < len(kept):
                    col, g = kept[i]
                else:
                    col, g = (kept[-1][0] if kept else first - 1) + (i - len(kept) + 1), ""
                e = expected[i] if i < len(expected) else ""
                if g != e:
                    diffs.append((f"{get_column_letter(col)}{hrow}", g, e))
            out[kind] = diffs
    finally:
        wb.close()
    return out


def template_alignment_problems(template_path: Path) -> List[str]:
    """Differences between the aligned TEMPLATE and the contract that are not
    allowed (empty list = the output headers, positions and bands match the
    TEMPLATE apart from the CHANNEL exclusion and whitespace)."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    problems: List[str] = []
    for kind, diffs in template_header_deviations(template_path).items():
        for cell, tmpl, contract in diffs:
            if contract == "(excluded)":
                continue
            if _ws(tmpl) != _ws(contract):
                problems.append(f"{kind} {cell}: TEMPLATE {tmpl!r} vs contract {contract!r}")
    wb = load_workbook(Path(template_path), read_only=False)
    try:
        for kind in ("premium", "claims", "outstanding"):
            sheet, hrow = TEMPLATE_SHEETS[kind]
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            _kept, _exc, colmap = _aligned_template(ws, kind)
            for col, want in _contract_bands(kind).items():
                tcols = [c for c, oc in colmap.items() if oc == col]
                got = _ws(ws.cell(hrow - 1, tcols[0]).value) if tcols else ""
                if got != _ws(want):
                    where = f"{get_column_letter(tcols[0])}{hrow - 1}" if tcols else "-"
                    problems.append(f"{kind} band {where}: TEMPLATE {got!r} vs contract {want!r}")
    finally:
        wb.close()
    return problems


def _contract_bands(kind: str) -> Dict[int, str]:
    """{output column: band label} on the row above the header row."""
    if kind == "premium":
        return dict(PREMIUM_BAND_LABELS)
    _row, first, _exp = _expected_headers(kind, DEFAULT_PROPORTION_MODE, True)
    return {first + CLAIMS_HEADERS.index("FROM"): "INSURANCE PERIOD"}


def validate_against_template(wb: Any, template_path: Optional[Path] = None, *,
                              claims_leading_blank: bool = True) -> List[str]:
    """Header row and band labels of every upload sheet of a written workbook
    against the aligned reference TEMPLATE (CHANNEL / SUB CHANNEL dropped,
    whitespace normalised). Empty list = same headers, positions and bands."""
    from openpyxl import load_workbook
    tpath = Path(template_path) if template_path else reference_template_path()
    twb = load_workbook(tpath, read_only=False)
    expected: Dict[str, Tuple[int, int, List[str], Dict[int, str]]] = {}
    try:
        for kind in ("premium", "claims", "outstanding"):
            sheet, hrow = TEMPLATE_SHEETS[kind]
            if sheet not in twb.sheetnames:
                continue
            tws = twb[sheet]
            kept, _exc, colmap = _aligned_template(tws, kind)
            bands = {}
            for tcol, ocol in colmap.items():
                v = _ws(tws.cell(hrow - 1, tcol).value)
                if v:
                    bands[ocol] = v
            _row, first, _exp = _expected_headers(kind, DEFAULT_PROPORTION_MODE, True)
            expected[kind] = (hrow, first, [_ws(t) for _c, t in kept], bands)
    finally:
        twb.close()
    problems: List[str] = []
    collapsed_names = {TEMPLATE_SHEETS[k][0]: k for k in ("premium", "claims", "outstanding")}
    for ws in wb.worksheets:
        kind, _label = _sheet_kind(ws.title, collapsed_names)
        if kind is None or kind not in expected:
            continue
        hrow, first, headers, bands = expected[kind]
        shift = 0 if kind == "premium" or claims_leading_blank else -(first - 1)
        got = [_ws(ws.cell(hrow, first + shift + i).value) for i in range(len(headers))]
        extra = _ws(ws.cell(hrow, first + shift + len(headers)).value)
        if got != headers or extra:
            diff = [f"{ws.cell(hrow, first + shift + i).coordinate}={g!r} (TEMPLATE {e!r})"
                    for i, (g, e) in enumerate(zip(got, headers)) if g != e]
            if extra:
                diff.append(f"extra header {extra!r} after the TEMPLATE's last column")
            problems.append(f"{ws.title!r} header row {hrow} vs TEMPLATE: " + "; ".join(diff[:6]))
        for col, text in bands.items():
            got_b = _ws(ws.cell(hrow - 1, col + shift).value)
            if got_b != text:
                problems.append(f"{ws.title!r} band {ws.cell(hrow - 1, col + shift).coordinate} "
                                f"is {got_b!r}, TEMPLATE {text!r}")
    return problems


def reference_template_path() -> Path:
    """The reference TEMPLATE.xlsx shipped with the backend (not the style seed)."""
    from src.domain.cre_cleaner.paths import TEMPLATES_DIR
    return TEMPLATES_DIR / "reference" / "TEMPLATE.xlsx"


__all__: Iterable[str] = (
    "TEMPLATE_DEVIATIONS_ALLOWED",
    "TEMPLATE_EXCLUDED_HEADERS",
    "reference_template_path",
    "template_alignment_problems",
    "template_header_deviations",
    "validate_against_template",
    "validate_upload_file",
    "validate_upload_workbook",
)
