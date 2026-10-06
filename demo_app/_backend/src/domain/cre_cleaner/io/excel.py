"""Read xls/xlsx, find headers, write cleaned workbook (upload / Bisola schema)."""
from __future__ import annotations

import contextlib
import contextvars
import math
import re
from copy import copy
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.domain.cre_cleaner.core.class_labels import (
    TYPE_CLAIMS,
    TYPE_OUTSTANDING,
    TYPE_PREMIUM,
    claims_class_hint,
    class_sheet_title,
    group_rows_by_class,
    ordered_class_labels,
    premium_class_hint,
)
from src.domain.cre_cleaner.config import (
    PREMIUM_COL_MAP,
    CLAIMS_COL_MAP,
    DEFAULT_PROPORTION_MODE,
    EXCEL_ERROR_STRINGS,
    FAC_CLASS_LABEL,
    PREMIUM_BAND_LABELS,
    TEMPLATE_SHEETS,
    premium_headers,
)
from src.domain.cre_cleaner.models import (
    PremiumRow,
    ClaimsRow,
    ExceptionRecord,
    SourceAuditRecord,
)
from src.domain.cre_cleaner.core.normalize import clean_text


# Request-scoped read cache (IMPL-20261002-01). One clean reads the same
# workbook several times (period inference, premium audit/load, claims
# audit/load); inside ``workbook_read_cache()`` each file is parsed once.
# Keyed by path + size + mtime + inode, and every hit returns fresh row lists,
# so a caller that edits its rows never changes what the next caller sees.
_READ_CACHE: contextvars.ContextVar[Optional[Dict[Any, Any]]] = contextvars.ContextVar(
    "cre_cleaner_workbook_read_cache", default=None,
)


@contextlib.contextmanager
def workbook_read_cache():
    """Parse each workbook at most once inside this block (nested blocks share it)."""
    if _READ_CACHE.get() is not None:
        yield
        return
    token = _READ_CACHE.set({})
    try:
        yield
    finally:
        _READ_CACHE.reset(token)


def _cache_key(kind: str, path: Path) -> Optional[Tuple[Any, ...]]:
    try:
        st = path.stat()
    except OSError:
        return None
    return (kind, str(path.resolve()), st.st_size, st.st_mtime_ns, st.st_ino)


def _copy_sheets(sheets: Dict[str, List[List[Any]]]) -> Dict[str, List[List[Any]]]:
    return {sn: [list(r) for r in rows] for sn, rows in sheets.items()}


def read_workbook_sheets(path: Path) -> Dict[str, List[List[Any]]]:
    """Return {sheet_name: rows as list of lists}. Supports xlsx and xls.

    Trailing rows that hold no value at all are not returned (xlsx sheets can
    declare ~65k styled but empty rows). Inside ``workbook_read_cache()`` a
    workbook is parsed once and later calls get a copy of the rows.
    """
    path = Path(path)
    cache = _READ_CACHE.get()
    key = _cache_key("sheets", path) if cache is not None else None
    if key is not None and key in cache:
        return _copy_sheets(cache[key])
    sheets = _read_workbook_sheets_uncached(path)
    if key is not None:
        cache[key] = sheets
        return _copy_sheets(sheets)
    return sheets


def _read_workbook_sheets_uncached(path: Path) -> Dict[str, List[List[Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _read_xls(path)
    try:
        return _read_xlsx(path)
    except Exception as e:
        # Misnamed OLE .xls or truncated OOXML often surface as BadZipFile
        name = type(e).__name__
        if name in {"BadZipFile", "BadZipfile"} or "not a zip file" in str(e).lower():
            try:
                return _read_xls(path)
            except Exception:
                raise RuntimeError(
                    f"Cannot read workbook {path.name!r}: file looks like Excel "
                    f"but is corrupt or truncated (not a valid .xlsx/.xls)."
                ) from e
        raise


def _xlsx_sheet_parts(zf: Any) -> List[Tuple[str, str, Any]]:
    """[(sheet name, worksheet part path, <sheet> element)] from the package."""
    import posixpath
    from xml.etree import ElementTree as ET

    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") or "" for r in rels}
    out = []
    for el in wb.iter():
        if not el.tag.endswith("}sheet"):
            continue
        target = targets.get(el.get(rel_id), "")
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
            posixpath.join("xl", target)
        )
        out.append((el.get("name"), part, el))
    return out


_ROW_R = re.compile(rb'\br="(\d+)"')


def _last_value_row(data: bytes) -> Optional[int]:
    """1-based index of the last <row> holding a value (<v> or inline <is>),
    0 when the sheet holds no value, None when it cannot be told safely."""
    if b":sheetData" in data:  # namespace-prefixed markup: do not guess
        return None
    end = data.rfind(b"</sheetData>")
    if end < 0:
        return 0 if b"<sheetData/>" in data else None
    pos = max(data.rfind(b"</v>", 0, end), data.rfind(b"</is>", 0, end))
    if pos < 0:
        return 0
    start = pos
    while True:
        start = data.rfind(b"<row", 0, start)
        if start < 0:
            return None
        nxt = data[start + 4:start + 5]
        if nxt in (b" ", b">", b"\t", b"\n", b"\r", b"/"):
            break
    tag_end = data.find(b">", start)
    m = _ROW_R.search(data, start, tag_end if tag_end > 0 else start + 200)
    return int(m.group(1)) if m else None


def _xlsx_last_value_rows(path: Path) -> Dict[str, Optional[int]]:
    """{sheet: last row with a value} read from the package XML (cheap)."""
    import zipfile

    out: Dict[str, Optional[int]] = {}
    with zipfile.ZipFile(path) as zf:
        for name, part, _el in _xlsx_sheet_parts(zf):
            try:
                out[name] = _last_value_row(zf.read(part))
            except KeyError:
                out[name] = None
    return out


def _read_xlsx(path: Path) -> Dict[str, List[List[Any]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        last = _xlsx_last_value_rows(path)
    except Exception:
        last = {}
    out: Dict[str, List[List[Any]]] = {}
    try:
        for sn in wb.sheetnames:
            ws = wb[sn]
            # Stop at the last row that holds a value: rows after it are all
            # None (styled phantom rows), so the rows returned are exactly the
            # previous ones without that empty tail. max_row=0 would mean
            # "no limit" to openpyxl, so an empty sheet is returned directly.
            lim = last.get(sn)
            if lim is not None and ws.max_row is not None:
                lim = min(lim, ws.max_row)
            if lim == 0:
                out[sn] = []
                continue
            rows = []
            for row in ws.iter_rows(values_only=True, max_row=lim):
                rows.append(list(row))
            out[sn] = rows
    finally:
        wb.close()
    return out


def _read_xls(path: Path) -> Dict[str, List[List[Any]]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    out: Dict[str, List[List[Any]]] = {}
    for sn in book.sheet_names():
        sh = book.sheet_by_name(sn)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                val = cell.value
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    # xlrd returns an int code (e.g. 15 for #VALUE!); never let it
                    # masquerade as a number downstream.
                    val = xlrd.error_text_from_code.get(val, "#N/A")
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        tup = xlrd.xldate_as_tuple(val, book.datemode)
                        from datetime import datetime
                        val = datetime(*tup[:6])
                    except Exception:
                        pass
                row.append(val)
            rows.append(row)
        out[sn] = rows
    return out


def sheet_names(path: Path) -> List[str]:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        import xlrd
        return xlrd.open_workbook(str(path)).sheet_names()
    try:
        wb = load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception as e:
        name = type(e).__name__
        if name in {"BadZipFile", "BadZipfile"} or "not a zip file" in str(e).lower():
            import xlrd
            try:
                return xlrd.open_workbook(str(path)).sheet_names()
            except Exception:
                raise RuntimeError(
                    f"Cannot read workbook {path.name!r}: file looks like Excel "
                    f"but is corrupt or truncated (not a valid .xlsx/.xls)."
                ) from e
        raise


_ROW_TAG = re.compile(rb"<(?:\w+:)?row\b([^>]*)>")
_COL_TAG = re.compile(rb"<(?:\w+:)?col\b([^>]*)>")
_ATTR_R = re.compile(rb'\br="(\d+)"')
_ATTR_MIN = re.compile(rb'\bmin="(\d+)"')
_ATTR_MAX = re.compile(rb'\bmax="(\d+)"')
_ATTR_HIDDEN = re.compile(rb'\bhidden="(?:1|true)"')


def _xlsx_visibility(path: Path) -> Dict[str, Dict[str, Any]]:
    """{sheet: {"state", "hidden_rows" (1-based), "hidden_cols" (1-based)}}.

    Read from the package XML: openpyxl's read-only mode (used for values)
    does not expose row/column dimensions, and a full load is much slower.
    """
    import posixpath
    import zipfile
    from xml.etree import ElementTree as ET

    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    out: Dict[str, Dict[str, Any]] = {}
    with zipfile.ZipFile(path) as zf:
        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target") or "" for r in rels}
        for el in wb.iter():
            if not el.tag.endswith("}sheet"):
                continue
            target = targets.get(el.get(rel_id), "")
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
                posixpath.join("xl", target)
            )
            try:
                data = zf.read(part)
            except KeyError:
                data = b""
            hidden_rows = set()
            hidden_cols: List[int] = []
            if b'hidden="' not in data:
                # No row/col can carry hidden="1|true": skip the tag scans
                # (same result, IMPL-20261002-03).
                out[el.get("name")] = {
                    "state": el.get("state") or "visible",
                    "hidden_rows": hidden_rows,
                    "hidden_cols": hidden_cols,
                }
                continue
            for m in _ROW_TAG.finditer(data):
                attrs = m.group(1)
                if _ATTR_HIDDEN.search(attrs):
                    r = _ATTR_R.search(attrs)
                    if r:
                        hidden_rows.add(int(r.group(1)))
            for m in _COL_TAG.finditer(data):
                attrs = m.group(1)
                if _ATTR_HIDDEN.search(attrs):
                    lo, hi = _ATTR_MIN.search(attrs), _ATTR_MAX.search(attrs)
                    if lo and hi:
                        hidden_cols.extend(range(int(lo.group(1)), min(int(hi.group(1)), 200) + 1))
            out[el.get("name")] = {
                "state": el.get("state") or "visible",
                "hidden_rows": hidden_rows,
                "hidden_cols": hidden_cols,
            }
    return out


def _xls_visibility(path: Path) -> Dict[str, Dict[str, Any]]:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    out: Dict[str, Dict[str, Any]] = {}
    try:
        for i, sn in enumerate(book.sheet_names()):
            sh = book.sheet_by_index(i)
            out[sn] = {
                "state": {0: "visible", 1: "hidden", 2: "veryHidden"}.get(sh.visibility, "visible"),
                "hidden_rows": {r + 1 for r, info in sh.rowinfo_map.items() if info.hidden},
                "hidden_cols": sorted(c + 1 for c, info in sh.colinfo_map.items() if info.hidden),
            }
    finally:
        book.release_resources()
    return out


def _read_visibility_uncached(path: Path) -> Dict[str, Dict[str, Any]]:
    try:
        return _xlsx_visibility(path)
    except Exception:
        return _xls_visibility(path)


def _copy_visibility(vis: Dict[str, Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    return {
        sn: {**v, "hidden_rows": set(v.get("hidden_rows") or ()),
             "hidden_cols": list(v.get("hidden_cols") or ())}
        for sn, v in vis.items()
    }


def _sheet_visibility(path: Path) -> Dict[str, Dict[str, Any]]:
    """Sheet state / hidden rows / hidden cols. Inside ``workbook_read_cache()``
    the scan runs once per file (the result, or the failure, is reused)."""
    cache = _READ_CACHE.get()
    key = _cache_key("visibility", path) if cache is not None else None
    if key is None:
        return _read_visibility_uncached(path)
    if key not in cache:
        try:
            cache[key] = ("ok", _read_visibility_uncached(path))
        except Exception as exc:  # re-raised on every call, like the uncached path
            cache[key] = ("error", exc)
    status, value = cache[key]
    if status == "error":
        raise value
    return _copy_visibility(value)


def read_source_workbook(
    path: Path,
) -> Tuple[Dict[str, List[List[Any]]], List[Tuple[str, str, str, int, str]]]:
    """Source reader for parsing: hidden sheets are dropped and hidden rows
    blanked (row positions kept so Excel row numbers stay right).

    Returns (sheets, notes); each note is (severity, reason, sheet, excel_row,
    detail) for the exceptions file. Visibility that cannot be read is itself
    reported rather than assumed visible silently.
    """
    path = Path(path)
    sheets = read_workbook_sheets(path)
    notes: List[Tuple[str, str, str, int, str]] = []
    try:
        vis = _sheet_visibility(path)
    except Exception as e:
        notes.append((
            "WARN", "visibility_unknown", "", 0,
            f"Could not read hidden row/sheet flags ({type(e).__name__}: {e}); "
            "hidden rows, if any, were read as normal rows",
        ))
        return sheets, notes

    out: Dict[str, List[List[Any]]] = {}
    for sn, rows in sheets.items():
        v = vis.get(sn) or {}
        if v.get("state", "visible") != "visible":
            # Hidden sheets stay unread; no WARN — the skip is normal for drafts.
            continue
        hidden = v.get("hidden_rows") or set()
        if hidden:
            rows = list(rows)
            for r1 in sorted(hidden):
                i = r1 - 1
                if not (0 <= i < len(rows)):
                    continue
                vals = [clean_text(c) for c in rows[i] if clean_text(c)]
                if vals:
                    notes.append((
                        "WARN", "hidden_row_skipped", sn, r1,
                        "Hidden in source, not read: " + " | ".join(vals[:8])[:300],
                    ))
                rows[i] = [None] * len(rows[i])
        if v.get("hidden_cols"):
            notes.append((
                "INFO", "hidden_columns_present", sn, 0,
                "Hidden columns still read (header-mapped): "
                + ", ".join(get_column_letter(c) for c in v["hidden_cols"][:30]),
            ))
        out[sn] = rows
    return out, notes


# ---------------------------------------------------------------------------
# Upload writer (fresh sheets; TEMPLATE supplies header styling only)
# ---------------------------------------------------------------------------
# Why fresh sheets: copying TEMPLATE sheets dragged in leftovers that broke the
# Continental validator — used range to Z1207/AL1000, a stray " " in S1, hidden
# column D (TEMPLATE's CHANNEL col, which becomes UNDERWRITING YEAR after the
# upload remap) and TEMPLATE date formats on col G (TOTAL SUM INSURED), which
# made date-aware readers (openpyxl/pandas/validator) render SI as #VALUE!.

PREMIUM_HEADER_ROW = 3
CLAIMS_HEADER_ROW = 4

_PREMIUM_AMOUNT_KEYS = {
    "TOTAL SUM INSURED", "GROSS PREMIUM", "RET SUM INSURED", "RET PREMIUM",
    "TREATY SUM INSURED", "TREATY PREMIUM", "FAC SUM INSURED", "FAC PREMIUM",
}
_PREMIUM_PCT_KEYS = {
    "MPL %", "RETENTION PROPORTION %", "TREATY PROPORTION %", "FACULTATIVE PROPORTION %",
}
_CLAIMS_AMOUNT_KEYS = {
    "TOTAL CLAIMS", "RET AMOUNT", "TREATY AMOUNT", "FAC AMOUNT",
}
_CLAIMS_PCT_KEYS = {"PPN RET %", "PPN TREATY %", "PPN FAC %"}

_AMOUNT_FMT = "#,##0.00"
# Premium proportions are stored as percent points (8.2 → "8.20"); claims PPN
# shares are fractions (amount/total → 0.062) and must use Excel % format so
# they display like Bisola gold (6.21%), not 0.06.
_PCT_FMT = "0.00"
_CLAIMS_PCT_FMT = "0.00%"
_DATE_FMT = "DD/MM/YYYY"

_PREMIUM_WIDTHS = [7, 22, 45, 12, 12, 12, 20, 9, 18, 12, 20, 18, 12, 20, 18, 12, 18, 16]
# TEMPLATE claims headers B..R (no SUM INSURED): S/NO … TO, TOTAL CLAIMS, … DETAILS
_CLAIMS_WIDTHS = [7, 40, 18, 22, 26, 14, 8, 12, 12, 18, 10, 18, 10, 18, 10, 16, 50]


class _Sink:
    """Collects writer-side data-quality exceptions (go to the sidecar)."""

    def __init__(self) -> None:
        self.records: List[ExceptionRecord] = []

    def add(self, severity: str, reason: str, audit=None, detail: str = "") -> None:
        fn = getattr(audit, "source_filename", "") if audit is not None else ""
        sh = getattr(audit, "source_sheet", "") if audit is not None else ""
        rw = getattr(audit, "source_row", 0) if audit is not None else 0
        self.records.append(ExceptionRecord(severity, reason, fn, sh, rw, detail))


def _is_excel_error(v: Any) -> bool:
    return isinstance(v, str) and v.strip().upper() in EXCEL_ERROR_STRINGS


def _sanitize(v: Any, *, key: str, sink: Optional[_Sink], audit=None, out_sheet: str = "") -> Any:
    """Never write Excel error literals, whitespace-only strings or NaN/inf."""
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        if sink is not None:
            sink.add("WARN", "non_finite_number_blanked", audit, f"{out_sheet}: {key}={v!r}")
        return None
    if isinstance(v, str):
        if _is_excel_error(v):
            if sink is not None:
                reason = (
                    "total_sum_insured_unrecoverable"
                    if key == "TOTAL SUM INSURED"
                    else "excel_error_value_blanked"
                )
                sink.add("WARN", reason, audit, f"{out_sheet}: {key}={v.strip()!r} left blank")
            return None
        s = v.strip()
        if not s:
            return None
        return s
    return v


def _template_styles(template_path: Optional[Path]) -> Dict[str, Dict[str, Any]]:
    """Pull header/band fonts, fills, borders, alignment from TEMPLATE (optional)."""
    styles: Dict[str, Dict[str, Any]] = {}
    if not template_path or not Path(template_path).exists():
        return styles
    try:
        twb = load_workbook(template_path)
    except Exception:
        return styles
    try:
        picks = {
            "premium_header": (TEMPLATE_SHEETS["premium"][0], "A3"),
            "premium_band": (TEMPLATE_SHEETS["premium"][0], "G2"),
            "claims_header": (TEMPLATE_SHEETS["claims"][0], "B4"),
            "claims_band": (TEMPLATE_SHEETS["claims"][0], "I3"),
        }
        for name, (sn, addr) in picks.items():
            if sn not in twb.sheetnames:
                continue
            c = twb[sn][addr]
            styles[name] = {
                "font": copy(c.font),
                "fill": copy(c.fill),
                "border": copy(c.border),
                "alignment": copy(c.alignment),
            }
    finally:
        twb.close()
    return styles


def _apply_style(cell, style: Optional[Dict[str, Any]], *, bold: bool = True) -> None:
    if style:
        cell.font = copy(style["font"])
        cell.fill = copy(style["fill"])
        cell.border = copy(style["border"])
        cell.alignment = copy(style["alignment"])
    else:
        cell.font = Font(bold=bold)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_title(ws: Worksheet, text: str, last_col: int, start_col: int = 1) -> None:
    c = ws.cell(1, start_col)
    c.value = text
    c.font = Font(bold=True, size=12)
    c.alignment = Alignment(horizontal="center", vertical="center")
    if last_col > start_col:
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=last_col)


def _set_widths(ws: Worksheet, widths: Sequence[float], start_col: int = 1) -> None:
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + i)].width = w


def enforce_upload_hygiene(
    ws: Worksheet, last_col: int, sink: Optional[_Sink] = None, first_col: int = 1,
) -> None:
    """Final guard: nothing outside first_col..last_col, no whitespace-only/error
    cells, nothing hidden. ``first_col=2`` keeps column A fully empty (claims /
    outstanding leading-blank layout)."""
    # Drop any cell beyond the last header column (incl. title/band rows) and,
    # for the leading-blank layout, anything in the reserved empty column(s).
    for (r, c) in list(ws._cells.keys()):
        if c > last_col or c < first_col:
            del ws._cells[(r, c)]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=last_col):
        for cell in row:
            v = cell.value
            if isinstance(v, str):
                if not v.strip():
                    cell.value = None
                elif _is_excel_error(v):
                    if sink is not None:
                        sink.add("WARN", "excel_error_value_blanked", None,
                                 f"{ws.title}!{cell.coordinate}={v.strip()!r}")
                    cell.value = None
    for key, dim in list(ws.column_dimensions.items()):
        dim.hidden = False
    for key, dim in list(ws.row_dimensions.items()):
        dim.hidden = False


def write_premium_rows(
    ws,
    rows: Sequence[PremiumRow],
    header_row: int = PREMIUM_HEADER_ROW,
    *,
    title: Optional[str] = None,
    proportion_mode: str = DEFAULT_PROPORTION_MODE,
    styles: Optional[Dict[str, Dict[str, Any]]] = None,
    sink: Optional[_Sink] = None,
) -> int:
    """Write an 18-column (A..R) upload premium sheet onto a blank worksheet."""
    styles = styles or {}
    headers = premium_headers(proportion_mode)
    last_col = len(headers)
    band_row = header_row - 1
    if title:
        _write_title(ws, title, last_col)
    for col, text in PREMIUM_BAND_LABELS.items():
        c = ws.cell(band_row, col)
        c.value = text
        _apply_style(c, styles.get("premium_band"))
    for start, end in ((5, 6), (10, 12), (13, 15), (16, 18)):
        ws.merge_cells(start_row=band_row, start_column=start, end_row=band_row, end_column=end)
    for c_i, h in enumerate(headers, start=1):
        c = ws.cell(header_row, c_i)
        c.value = h
        _apply_style(c, styles.get("premium_header"))
    _set_widths(ws, _PREMIUM_WIDTHS)

    for i, prow in enumerate(rows, start=1):
        r = header_row + i
        vals = prow.to_template_values()
        sno = ws.cell(r, PREMIUM_COL_MAP["S/NO"])
        sno.value = i
        sno.number_format = "0"
        for key, col in PREMIUM_COL_MAP.items():
            if key == "S/NO":
                continue
            v = _sanitize(vals.get(key), key=key, sink=sink, audit=prow.audit, out_sheet=ws.title)
            if key == "TOTAL SUM INSURED" and v is None and sink is not None:
                sink.add("INFO", "total_sum_insured_blank", prow.audit,
                         f"{ws.title}: policy={prow.policy_no!r} TSI blank in source")
            cell = ws.cell(r, col)
            cell.value = v
            # Explicit format on EVERY data cell (never inherit a date format).
            if key in ("FROM", "TO"):
                cell.number_format = _DATE_FMT
            elif key in ("POLICY NO.", "NAME OF INSURED"):
                cell.number_format = "@"
            elif key in _PREMIUM_AMOUNT_KEYS:
                cell.number_format = _AMOUNT_FMT
            elif key in _PREMIUM_PCT_KEYS:
                cell.number_format = _PCT_FMT
            elif key == "UNDERWRITING YEAR":
                cell.number_format = "0"
            else:
                cell.number_format = "General"
    enforce_upload_hygiene(ws, last_col, sink)
    return len(rows)


def claims_col_map(leading_blank: bool = True) -> Dict[str, int]:
    """CLAIMS_COL_MAP holds TEMPLATE positions (S/NO. in col B).

    Default upload layout (TEMPLATE.xlsx): column A fully empty, headers B..R.
    ``leading_blank=False`` shifts everything to start at col A."""
    shift = 0 if leading_blank else -(min(CLAIMS_COL_MAP.values()) - 1)
    return {k: v + shift for k, v in CLAIMS_COL_MAP.items()}


def write_claims_rows(
    ws,
    rows: Sequence[ClaimsRow],
    header_row: int = CLAIMS_HEADER_ROW,
    *,
    force_class_label: Optional[str] = None,
    title: Optional[str] = None,
    leading_blank: bool = True,
    styles: Optional[Dict[str, Dict[str, Any]]] = None,
    sink: Optional[_Sink] = None,
) -> int:
    """Write a claims/outstanding sheet matching TEMPLATE.xlsx.

    Default (``leading_blank=True``): column A fully empty, title in B1, headers
    B..R. ``leading_blank=False``: headers start at column A."""
    styles = styles or {}
    cmap = claims_col_map(leading_blank)
    first_col = min(cmap.values())
    last_col = max(cmap.values())
    period_row = header_row - 1
    if title:
        _write_title(ws, title, last_col, start_col=first_col)
    pc = ws.cell(period_row, cmap["FROM"])
    pc.value = "INSURANCE PERIOD"
    _apply_style(pc, styles.get("claims_band"))
    ws.merge_cells(start_row=period_row, start_column=cmap["FROM"],
                   end_row=period_row, end_column=cmap["TO"])
    for key, col in cmap.items():
        c = ws.cell(header_row, col)
        c.value = key
        _apply_style(c, styles.get("claims_header"))
    _set_widths(ws, _CLAIMS_WIDTHS, start_col=first_col)

    for i, crow in enumerate(rows, start=1):
        r = header_row + i
        vals = crow.to_template_values()
        if force_class_label:
            vals = dict(vals)
            vals["CLASS"] = force_class_label
        sno = ws.cell(r, cmap["S/NO."])
        sno.value = i
        sno.number_format = "0"
        for key, col in cmap.items():
            if key == "S/NO.":
                continue
            v = _sanitize(vals.get(key), key=key, sink=sink, audit=crow.audit, out_sheet=ws.title)
            cell = ws.cell(r, col)
            cell.value = v
            if key in ("DATE OF LOSS (day-mth-year)", "FROM", "TO"):
                cell.number_format = _DATE_FMT
            elif key in ("POLICY NO.", "CLAIM NO", "INSURED", "CLASS", "DETAILS OF LOSS"):
                cell.number_format = "@"
            elif key in _CLAIMS_AMOUNT_KEYS:
                cell.number_format = _AMOUNT_FMT
            elif key in _CLAIMS_PCT_KEYS:
                cell.number_format = _CLAIMS_PCT_FMT
            elif key == "UW YR":
                cell.number_format = "0"
            else:
                cell.number_format = "General"
    enforce_upload_hygiene(ws, last_col, sink, first_col=first_col)
    return len(rows)


def _ensure_sheet(wb, title: str):
    if title in wb.sheetnames:
        return wb[title]
    return wb.create_sheet(title)


def _set_sheet_banner(ws: Worksheet, text: str, cell_addr: str = "A1") -> None:
    """Overwrite title/banner cell when present (class-split sheets)."""
    try:
        ws[cell_addr] = text
    except Exception:
        pass


def write_summary_sheet(wb, summary: dict) -> None:
    ws = _ensure_sheet(wb, "SUMMARY")
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    ws["A1"] = "CRE CLEANER — RUN SUMMARY"
    ws["A2"] = "Cedant"
    ws["B2"] = summary.get("cedant")
    ws["A3"] = "Broker"
    ws["B3"] = summary.get("broker")
    ws["A4"] = "Year"
    ws["B4"] = summary.get("year")
    ws["A5"] = "Quarter"
    ws["B5"] = summary.get("quarter")
    ws["A6"] = "Generated (Africa/Lagos)"
    ws["B6"] = summary.get("generated_at")
    ws["A7"] = "Output layout"
    ws["B7"] = summary.get("output_layout", "class-split")
    ws["A8"] = "Currency"
    ws["B8"] = summary.get("currency", "NGN")

    ws["A9"] = "Metric"
    ws["B9"] = "Value"
    splits = summary.get("split_checks") or {}
    dates = summary.get("date_checks") or {}
    metrics = [
        ("Adapter status", summary.get("adapter_status")),
        ("Premium source files", summary.get("premium_files")),
        ("Claims source files", summary.get("claims_files")),
        ("Premium rows written", summary.get("premium_rows")),
        ("Paid claims rows written", summary.get("claims_rows")),
        ("Outstanding rows written", summary.get("outstanding_rows")),
        ("Exception rows", summary.get("exception_rows")),
        ("Premium gross premium sum", summary.get("premium_gross_sum")),
        ("Claims total sum", summary.get("claims_total_sum")),
        ("Outstanding total sum", summary.get("outstanding_total_sum")),
        ("Outstanding supplied", summary.get("outstanding_supplied")),
    ]
    for label, key in (("premium", "PREMIUM"), ("claims", "CLAIMS"), ("outstanding", "OUTSTANDING")):
        s = splits.get(label)
        if s:
            metrics.append((
                f"Split check {key} (OK / mismatch / not checkable)",
                f"{s.get('ok', 0)} / {s.get('mismatch', 0)} / {s.get('not_checkable', 0)}",
            ))
    if dates:
        metrics.append((
            "Date check flags",
            "; ".join(f"{k}: {v}" for k, v in sorted(dates.items())) or "none",
        ))
    metrics.append(("Notes", summary.get("notes")))
    row_i = 10
    for k, v in metrics:
        ws.cell(row_i, 1).value = k
        ws.cell(row_i, 2).value = v
        row_i += 1

    by_class = summary.get("by_class") or {}
    if by_class:
        row_i += 1
        ws.cell(row_i, 1).value = "Class"
        ws.cell(row_i, 2).value = "Premium rows"
        ws.cell(row_i, 3).value = "Claims rows"
        ws.cell(row_i, 4).value = "Outstanding rows"
        row_i += 1
        for lab, counts in by_class.items():
            ws.cell(row_i, 1).value = lab
            ws.cell(row_i, 2).value = counts.get("premium", 0)
            ws.cell(row_i, 3).value = counts.get("claims", 0)
            ws.cell(row_i, 4).value = counts.get("outstanding", 0)
            row_i += 1

    recon = summary.get("reconciliation") or []
    if recon:
        row_i += 1
        ws.cell(row_i, 1).value = "SOURCE RECONCILIATION (output vs source, same currency only)"
        ws.cell(row_i, 1).font = Font(bold=True)
        row_i += 1
        headers = [
            "Level", "Period", "Bordereau", "Metric", "Source total row",
            "Source row sum", "Output", "Variance", "Status",
            "Printed total - row sum",
        ]
        for c, h in enumerate(headers, 1):
            ws.cell(row_i, c).value = h
            ws.cell(row_i, c).font = Font(bold=True)
        row_i += 1
        for rec in recon:
            vals = [
                rec.get("level"), rec.get("period"), rec.get("bordereau"), rec.get("metric"),
                rec.get("source_total_row"), rec.get("source_row_sum"), rec.get("output"),
                rec.get("variance"), rec.get("status"), rec.get("printed_total_variance"),
            ]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row_i, c)
                cell.value = v
                if c in (5, 6, 7, 8, 10) and isinstance(v, float):
                    cell.number_format = _AMOUNT_FMT
            row_i += 1


def _write_tabular_workbook(
    path: Path,
    sheet_title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c).value = h
    for i, rec in enumerate(rows, start=2):
        for c, v in enumerate(rec, 1):
            ws.cell(i, c).value = v
    wb.save(path)
    wb.close()
    return path


def write_exceptions_sidecar(
    path: Path,
    records: Sequence[ExceptionRecord],
) -> Path:
    headers = [
        "Severity", "Reason", "Source Filename", "Source Sheet",
        "Source Row", "Detail",
    ]
    return _write_tabular_workbook(
        path, "EXCEPTIONS", headers, [r.as_row() for r in records]
    )


def write_source_audit_sidecar(
    path: Path,
    records: Sequence[SourceAuditRecord],
) -> Path:
    return _write_tabular_workbook(
        path, "SOURCE AUDIT", SourceAuditRecord.AUDIT_HEADERS, [r.as_row() for r in records]
    )


def write_source_audit_sheet(wb, records: Sequence[SourceAuditRecord]) -> None:
    """Optional: embed SOURCE AUDIT in the main workbook (--include-audit-sheets)."""
    ws = _ensure_sheet(wb, "SOURCE AUDIT")
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    for c, h in enumerate(SourceAuditRecord.AUDIT_HEADERS, 1):
        ws.cell(1, c).value = h
    for i, rec in enumerate(records, start=2):
        for c, v in enumerate(rec.as_row(), 1):
            ws.cell(i, c).value = v


def write_exceptions_sheet(wb, records: Sequence[ExceptionRecord]) -> None:
    """Optional: embed EXCEPTIONS in the main workbook (--include-audit-sheets)."""
    ws = _ensure_sheet(wb, "EXCEPTIONS")
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    headers = [
        "Severity", "Reason", "Source Filename", "Source Sheet",
        "Source Row", "Detail",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c).value = h
    ws.cell(1, 8).value = "(DATA QUALITY LOG)"
    for i, rec in enumerate(records, start=2):
        for c, v in enumerate(rec.as_row(), 1):
            ws.cell(i, c).value = v


def _remove_sheet_if_present(wb, title: str) -> None:
    if title in wb.sheetnames:
        wb.remove(wb[title])


def _filter_fac_labels(labels, include_fac: bool):
    """Facultative is never its own upload sheet (IMPL-20261006-03 B3);
    ``include_fac`` is kept for call compatibility and ignored."""
    return [lab for lab in labels if lab != FAC_CLASS_LABEL]


def _write_class_split_sheets(
    wb,
    premium_rows: Sequence[PremiumRow],
    claims_rows: Sequence[ClaimsRow],
    outstanding_rows: Sequence[ClaimsRow],
    *,
    year: Optional[int] = None,
    quarter: Optional[int] = None,
    include_fac: bool = False,
    proportion_mode: str = DEFAULT_PROPORTION_MODE,
    claims_leading_blank: bool = True,
    styles: Optional[Dict[str, Dict[str, Any]]] = None,
    sink: Optional[_Sink] = None,
    class_map: Any = None,
) -> Dict[str, int]:
    """Create Bisola-style per-class sheets; omit empty class×type combos.

    Returns {sheet_title: row_count}.
    """
    prem_by = group_rows_by_class(premium_rows, premium_class_hint, class_map)
    paid_by = group_rows_by_class(claims_rows, claims_class_hint, class_map)
    ost_by = group_rows_by_class(outstanding_rows, claims_class_hint, class_map)

    all_labels = _filter_fac_labels(
        ordered_class_labels(set(prem_by) | set(paid_by) | set(ost_by)),
        include_fac,
    )

    inventory: Dict[str, int] = {}
    q_bit = f" — Q{quarter} {year}" if year and quarter else ""

    for lab in all_labels:
        rows_p = prem_by.get(lab) or []
        if rows_p:
            title = class_sheet_title(lab, TYPE_PREMIUM)
            ws = wb.create_sheet(title)
            inventory[title] = write_premium_rows(
                ws, rows_p, PREMIUM_HEADER_ROW,
                title=f"{lab.upper()} — PREMIUM BORDEREAU{q_bit}",
                proportion_mode=proportion_mode, styles=styles, sink=sink,
            )

        rows_c = paid_by.get(lab) or []
        if rows_c:
            title = class_sheet_title(lab, TYPE_CLAIMS)
            ws = wb.create_sheet(title)
            inventory[title] = write_claims_rows(
                ws, rows_c, CLAIMS_HEADER_ROW, force_class_label=lab,
                title=f"{lab.upper()} — CLAIMS BORDEREAU{q_bit}",
                leading_blank=claims_leading_blank, styles=styles, sink=sink,
            )

        rows_o = ost_by.get(lab) or []
        if rows_o:
            title = class_sheet_title(lab, TYPE_OUTSTANDING)
            ws = wb.create_sheet(title)
            inventory[title] = write_claims_rows(
                ws, rows_o, CLAIMS_HEADER_ROW, force_class_label=lab,
                title=f"{lab.upper()} — OUTSTANDING LOSS BORDEREAU{q_bit}",
                leading_blank=claims_leading_blank, styles=styles, sink=sink,
            )

    return inventory


def _write_collapsed_sheets(
    wb,
    premium_rows: Sequence[PremiumRow],
    claims_rows: Sequence[ClaimsRow],
    outstanding_rows: Sequence[ClaimsRow],
    *,
    proportion_mode: str = DEFAULT_PROPORTION_MODE,
    claims_leading_blank: bool = True,
    styles: Optional[Dict[str, Dict[str, Any]]] = None,
    sink: Optional[_Sink] = None,
) -> Dict[str, int]:
    """Legacy single PREMIUM / CLAIMS / OUTSTANDING sheets."""
    prem_name = TEMPLATE_SHEETS["premium"][0]
    clm_name = TEMPLATE_SHEETS["claims"][0]
    ost_name = TEMPLATE_SHEETS["outstanding"][0]
    inventory: Dict[str, int] = {}
    inventory[prem_name] = write_premium_rows(
        wb.create_sheet(prem_name), premium_rows, PREMIUM_HEADER_ROW,
        title=prem_name, proportion_mode=proportion_mode, styles=styles, sink=sink,
    )
    inventory[clm_name] = write_claims_rows(
        wb.create_sheet(clm_name), claims_rows, CLAIMS_HEADER_ROW,
        title=clm_name, leading_blank=claims_leading_blank, styles=styles, sink=sink,
    )
    inventory[ost_name] = write_claims_rows(
        wb.create_sheet(ost_name), outstanding_rows, CLAIMS_HEADER_ROW,
        title=ost_name, leading_blank=claims_leading_blank, styles=styles, sink=sink,
    )
    return inventory


def _reorder_sheets(wb, preferred: List[str]) -> None:
    """Move sheets so preferred titles come first (in order), rest follow."""
    existing = list(wb.sheetnames)
    front = [t for t in preferred if t in existing]
    rest = [t for t in existing if t not in front]
    target = front + rest
    for idx, title in enumerate(target):
        current = wb.sheetnames.index(title)
        offset = idx - current
        if offset:
            wb.move_sheet(wb[title], offset=offset)


def _summary_last_col(ws: Worksheet) -> int:
    last = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and not (isinstance(c.value, str) and not c.value.strip()):
                last = max(last, c.column)
    return max(last, 1)


def write_output_workbook(
    template_path: Path,
    output_path: Path,
    premium_rows: Sequence[PremiumRow],
    claims_rows: Sequence[ClaimsRow],
    outstanding_rows: Sequence[ClaimsRow],
    exceptions: Sequence[ExceptionRecord],
    source_audit: Sequence[SourceAuditRecord],
    summary: dict,
    *,
    collapsed: bool = False,
    include_audit_sheets: bool = False,
    include_fac: bool = False,
    proportion_mode: str = DEFAULT_PROPORTION_MODE,
    claims_leading_blank: bool = True,
    class_map: Any = None,
) -> Tuple[Path, Path, Path]:
    """Write upload-ready workbook + sidecar exception/audit files.

    Sheets are built fresh (TEMPLATE is a style seed only) so no TEMPLATE
    leftovers (hidden cols, stray cells, date formats, wide used range) leak in.
    Returns (cleaned_xlsx, exceptions_sidecar, source_audit_sidecar).
    """
    premium_headers(proportion_mode)  # validate mode early
    from src.domain.cre_cleaner.core.class_labels import (
        claims_class_hint,
        format_unapproved_classes,
        premium_class_hint,
        unapproved_class_rows,
    )

    blocked = (
        unapproved_class_rows(premium_rows, premium_class_hint, class_map)
        + unapproved_class_rows(claims_rows, claims_class_hint, class_map)
        + unapproved_class_rows(outstanding_rows, claims_class_hint, class_map)
    )
    if blocked:
        raise ValueError(format_unapproved_classes(blocked))
    template_path = Path(template_path) if template_path else None
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    styles = _template_styles(template_path)
    sink = _Sink()
    wb = Workbook()
    wb.remove(wb.active)

    if collapsed:
        inventory = _write_collapsed_sheets(
            wb, premium_rows, claims_rows, outstanding_rows,
            proportion_mode=proportion_mode,
            claims_leading_blank=claims_leading_blank,
            styles=styles, sink=sink,
        )
        layout = "collapsed"
    else:
        inventory = _write_class_split_sheets(
            wb,
            premium_rows,
            claims_rows,
            outstanding_rows,
            year=summary.get("year"),
            quarter=summary.get("quarter"),
            include_fac=include_fac,
            proportion_mode=proportion_mode,
            claims_leading_blank=claims_leading_blank,
            styles=styles,
            sink=sink,
            class_map=class_map,
        )
        layout = "class-split"

    all_exceptions = list(exceptions) + sink.records
    try:
        summary["output_layout"] = layout
        summary["sheet_inventory"] = inventory
        summary["proportion_headers"] = proportion_mode
        summary["exception_rows"] = len(all_exceptions)
    except TypeError:
        summary = dict(summary)
        summary["output_layout"] = layout
        summary["sheet_inventory"] = inventory
        summary["proportion_headers"] = proportion_mode
        summary["exception_rows"] = len(all_exceptions)

    write_summary_sheet(wb, summary)
    ws_sum = wb["SUMMARY"]
    enforce_upload_hygiene(ws_sum, _summary_last_col(ws_sum), sink)

    # Sidecars always (debugging without poisoning the upload workbook)
    stem = output_path.with_suffix("")
    exc_path = Path(str(stem) + "_exceptions.xlsx")
    audit_path = Path(str(stem) + "_source_audit.xlsx")
    all_exceptions = list(exceptions) + sink.records
    write_exceptions_sidecar(exc_path, all_exceptions)
    write_source_audit_sidecar(audit_path, source_audit)

    preferred = ["SUMMARY"] + list(inventory.keys())
    if include_audit_sheets:
        write_source_audit_sheet(wb, source_audit)
        write_exceptions_sheet(wb, all_exceptions)
        preferred = preferred + ["SOURCE AUDIT", "EXCEPTIONS"]

    _reorder_sheets(wb, preferred)
    wb.active = 0

    # Upload-structure check on the finished workbook (IMPL-20261006-03 B5).
    # A problem is an ERROR the service will not publish.
    from src.domain.cre_cleaner.io.upload_validator import validate_upload_workbook
    for problem in validate_upload_workbook(
        wb, proportion_mode=proportion_mode,
        claims_leading_blank=claims_leading_blank,
        allow_audit_sheets=include_audit_sheets,
    ):
        sink.add("ERROR", "upload_structure_invalid", None, problem)

    wb.save(output_path)
    wb.close()
    # expose writer-side exceptions to callers (pipeline appends them)
    write_output_workbook.last_writer_exceptions = list(sink.records)
    return output_path, exc_path, audit_path
