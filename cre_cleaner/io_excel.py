"""Read xls/xlsx, find headers, write cleaned workbook from TEMPLATE."""
from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from cre_cleaner.class_labels import (
    TYPE_CLAIMS,
    TYPE_OUTSTANDING,
    TYPE_PREMIUM,
    claims_class_hint,
    class_sheet_title,
    group_rows_by_class,
    ordered_class_labels,
    premium_class_hint,
)
from cre_cleaner.config import (
    PREMIUM_COL_MAP,
    CLAIMS_COL_MAP,
    TEMPLATE_SHEETS,
)
from cre_cleaner.models import (
    PremiumRow,
    ClaimsRow,
    ExceptionRecord,
    SourceAuditRecord,
)


def read_workbook_sheets(path: Path) -> Dict[str, List[List[Any]]]:
    """Return {sheet_name: rows as list of lists}. Supports xlsx and xls."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _read_xls(path)
    return _read_xlsx(path)


def _read_xlsx(path: Path) -> Dict[str, List[List[Any]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    out: Dict[str, List[List[Any]]] = {}
    try:
        for sn in wb.sheetnames:
            ws = wb[sn]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
            out[sn] = rows
    finally:
        wb.close()
    return out


def _read_xls(path: Path) -> Dict[str, List[List[Any]]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    out: Dict[str, List[List[Any]]] = {}
    for sn in book.sheet_names():
        sh = book.sheet_by_name(sn)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                val = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        tup = xlrd.xldate_as_tuple(val, book.datemode)
                        from datetime import datetime
                        val = datetime(*tup[:6])
                    except Exception:
                        pass
                row.append(val)
            rows.append(row)
        out[sn] = rows
    return out


def sheet_names(path: Path) -> List[str]:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        import xlrd
        return xlrd.open_workbook(str(path)).sheet_names()
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def clear_data_rows(ws, header_row: int, start_col: int, end_col: int) -> None:
    """Clear existing data below header (keep formatting on row header_row+1 as template)."""
    max_r = ws.max_row or header_row
    for r in range(header_row + 1, max_r + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(r, c).value = None


def write_premium_rows(ws, rows: Sequence[PremiumRow], header_row: int = 3) -> int:
    clear_data_rows(ws, header_row, 1, 20)
    date_fmt = "DD/MM/YYYY"
    for i, prow in enumerate(rows, start=1):
        r = header_row + i
        vals = prow.to_template_values()
        ws.cell(r, PREMIUM_COL_MAP["S/NO"]).value = i
        for key, col in PREMIUM_COL_MAP.items():
            if key == "S/NO":
                continue
            v = vals.get(key)
            cell = ws.cell(r, col)
            cell.value = v
            if key in ("FROM", "TO") and v is not None:
                cell.number_format = date_fmt
            if key == "POLICY NO." and v is not None:
                cell.number_format = "@"
    return len(rows)


def write_claims_rows(
    ws,
    rows: Sequence[ClaimsRow],
    header_row: int = 4,
    *,
    force_class_label: Optional[str] = None,
) -> int:
    clear_data_rows(ws, header_row, 2, 18)
    date_fmt = "DD/MM/YYYY"
    for i, crow in enumerate(rows, start=1):
        r = header_row + i
        vals = crow.to_template_values()
        if force_class_label:
            vals = dict(vals)
            vals["CLASS"] = force_class_label
        ws.cell(r, CLAIMS_COL_MAP["S/NO."]).value = i
        for key, col in CLAIMS_COL_MAP.items():
            if key == "S/NO.":
                continue
            v = vals.get(key)
            cell = ws.cell(r, col)
            cell.value = v
            if key in ("DATE OF LOSS (day-mth-year)", "FROM", "TO") and v is not None:
                cell.number_format = date_fmt
            if key in ("POLICY NO.", "CLAIM NO") and v is not None:
                cell.number_format = "@"
    return len(rows)


def _ensure_sheet(wb, title: str):
    if title in wb.sheetnames:
        return wb[title]
    return wb.create_sheet(title)


def _set_sheet_banner(ws: Worksheet, text: str, cell_addr: str = "A1") -> None:
    """Overwrite title/banner cell when present (class-split sheets)."""
    try:
        ws[cell_addr] = text
    except Exception:
        pass


def write_summary_sheet(wb, summary: dict) -> None:
    ws = _ensure_sheet(wb, "SUMMARY")
    # clear
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    ws["A1"] = "CRE CLEANER — RUN SUMMARY"
    ws["A2"] = "Cedant"
    ws["B2"] = summary.get("cedant")
    ws["A3"] = "Broker"
    ws["B3"] = summary.get("broker")
    ws["A4"] = "Year"
    ws["B4"] = summary.get("year")
    ws["A5"] = "Quarter"
    ws["B5"] = summary.get("quarter")
    ws["A6"] = "Generated (Africa/Lagos)"
    ws["B6"] = summary.get("generated_at")
    ws["A7"] = "Output layout"
    ws["B7"] = summary.get("output_layout", "class-split")

    ws["A9"] = "Metric"
    ws["B9"] = "Value"
    metrics = [
        ("Premium source files", summary.get("premium_files")),
        ("Claims source files", summary.get("claims_files")),
        ("Premium rows written", summary.get("premium_rows")),
        ("Paid claims rows written", summary.get("claims_rows")),
        ("Outstanding rows written", summary.get("outstanding_rows")),
        ("Exception rows", summary.get("exception_rows")),
        ("Premium gross premium sum", summary.get("premium_gross_sum")),
        ("Claims total sum", summary.get("claims_total_sum")),
        ("Outstanding total sum", summary.get("outstanding_total_sum")),
        ("Outstanding supplied", summary.get("outstanding_supplied")),
        ("Notes", summary.get("notes")),
    ]
    row_i = 10
    for k, v in metrics:
        ws.cell(row_i, 1).value = k
        ws.cell(row_i, 2).value = v
        row_i += 1

    by_class = summary.get("by_class") or {}
    if by_class:
        row_i += 1
        ws.cell(row_i, 1).value = "Class"
        ws.cell(row_i, 2).value = "Premium rows"
        ws.cell(row_i, 3).value = "Claims rows"
        ws.cell(row_i, 4).value = "Outstanding rows"
        row_i += 1
        for lab, counts in by_class.items():
            ws.cell(row_i, 1).value = lab
            ws.cell(row_i, 2).value = counts.get("premium", 0)
            ws.cell(row_i, 3).value = counts.get("claims", 0)
            ws.cell(row_i, 4).value = counts.get("outstanding", 0)
            row_i += 1


def write_source_audit_sheet(wb, records: Sequence[SourceAuditRecord]) -> None:
    ws = _ensure_sheet(wb, "SOURCE AUDIT")
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    headers = [
        "Source Filename", "Source Sheet", "Sheet Type", "Header Row",
        "Rows Read", "Rows Kept", "Rows Skipped", "Source Month", "Notes",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c).value = h
    for i, rec in enumerate(records, start=2):
        for c, v in enumerate(rec.as_row(), 1):
            ws.cell(i, c).value = v


def write_exceptions_sheet(wb, records: Sequence[ExceptionRecord]) -> None:
    ws = _ensure_sheet(wb, "EXCEPTIONS")
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    headers = [
        "Severity", "Reason", "Source Filename", "Source Sheet",
        "Source Row", "Detail",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c).value = h
    # Also alias title note
    ws.cell(1, 8).value = "(DATA QUALITY LOG)"
    for i, rec in enumerate(records, start=2):
        for c, v in enumerate(rec.as_row(), 1):
            ws.cell(i, c).value = v


def _copy_template_sheet(wb, source_title: str, new_title: str) -> Worksheet:
    """Copy a TEMPLATE sheet (structure/headers/styles) under a new name."""
    if source_title not in wb.sheetnames:
        raise KeyError(f"TEMPLATE missing sheet {source_title!r}")
    src = wb[source_title]
    # openpyxl copy_worksheet keeps a "Copy of …" name; rename after
    ws = wb.copy_worksheet(src)
    ws.title = new_title
    return ws


def _remove_sheet_if_present(wb, title: str) -> None:
    if title in wb.sheetnames:
        wb.remove(wb[title])


def _write_class_split_sheets(
    wb,
    premium_rows: Sequence[PremiumRow],
    claims_rows: Sequence[ClaimsRow],
    outstanding_rows: Sequence[ClaimsRow],
    *,
    year: Optional[int] = None,
    quarter: Optional[int] = None,
) -> Dict[str, int]:
    """Create Bisola-style per-class sheets; omit empty class×type combos.

    Returns {sheet_title: row_count}.
    """
    prem_name, prem_hdr = TEMPLATE_SHEETS["premium"]
    clm_name, clm_hdr = TEMPLATE_SHEETS["claims"]
    ost_name, ost_hdr = TEMPLATE_SHEETS["outstanding"]

    prem_by = group_rows_by_class(premium_rows, premium_class_hint)
    paid_by = group_rows_by_class(claims_rows, claims_class_hint)
    ost_by = group_rows_by_class(outstanding_rows, claims_class_hint)

    all_labels = ordered_class_labels(
        set(prem_by) | set(paid_by) | set(ost_by)
    )

    inventory: Dict[str, int] = {}
    q_bit = f" — Q{quarter} {year}" if year and quarter else ""

    for lab in all_labels:
        # PREMIUM
        rows_p = prem_by.get(lab) or []
        if rows_p:
            title = class_sheet_title(lab, TYPE_PREMIUM)
            ws = _copy_template_sheet(wb, prem_name, title)
            _set_sheet_banner(ws, f"{lab.upper()} — PREMIUM BORDEREAU{q_bit}")
            n = write_premium_rows(ws, rows_p, prem_hdr)
            inventory[title] = n

        # CLAIMS (paid)
        rows_c = paid_by.get(lab) or []
        if rows_c:
            title = class_sheet_title(lab, TYPE_CLAIMS)
            ws = _copy_template_sheet(wb, clm_name, title)
            _set_sheet_banner(ws, f"{lab.upper()} — CLAIMS BORDEREAU{q_bit}")
            n = write_claims_rows(ws, rows_c, clm_hdr, force_class_label=lab)
            inventory[title] = n

        # OUTSTANDING
        rows_o = ost_by.get(lab) or []
        if rows_o:
            title = class_sheet_title(lab, TYPE_OUTSTANDING)
            # Outstanding TEMPLATE header row may differ; use configured hdr
            ws = _copy_template_sheet(wb, ost_name, title)
            _set_sheet_banner(ws, f"{lab.upper()} — OUTSTANDING LOSS BORDEREAU{q_bit}")
            n = write_claims_rows(ws, rows_o, ost_hdr, force_class_label=lab)
            inventory[title] = n

    # Drop the original collapsed TEMPLATE sheets (data now on class sheets)
    for name in (prem_name, clm_name, ost_name):
        _remove_sheet_if_present(wb, name)

    return inventory


def _write_collapsed_sheets(
    wb,
    premium_rows: Sequence[PremiumRow],
    claims_rows: Sequence[ClaimsRow],
    outstanding_rows: Sequence[ClaimsRow],
) -> Dict[str, int]:
    """Legacy single PREMIUM / CLAIMS / OUTSTANDING sheets."""
    prem_name, prem_hdr = TEMPLATE_SHEETS["premium"]
    clm_name, clm_hdr = TEMPLATE_SHEETS["claims"]
    ost_name, ost_hdr = TEMPLATE_SHEETS["outstanding"]
    inventory: Dict[str, int] = {}
    if prem_name in wb.sheetnames:
        inventory[prem_name] = write_premium_rows(wb[prem_name], premium_rows, prem_hdr)
    if clm_name in wb.sheetnames:
        inventory[clm_name] = write_claims_rows(wb[clm_name], claims_rows, clm_hdr)
    if ost_name in wb.sheetnames:
        inventory[ost_name] = write_claims_rows(wb[ost_name], outstanding_rows, ost_hdr)
    return inventory


def _reorder_sheets(wb, preferred: List[str]) -> None:
    """Move sheets so preferred titles come first (in order), rest follow."""
    # openpyxl: wb.move_sheet(sheet, offset=…)
    existing = list(wb.sheetnames)
    # Build target order
    front = [t for t in preferred if t in existing]
    rest = [t for t in existing if t not in front]
    target = front + rest
    for idx, title in enumerate(target):
        current = wb.sheetnames.index(title)
        offset = idx - current
        if offset:
            wb.move_sheet(wb[title], offset=offset)


def write_output_workbook(
    template_path: Path,
    output_path: Path,
    premium_rows: Sequence[PremiumRow],
    claims_rows: Sequence[ClaimsRow],
    outstanding_rows: Sequence[ClaimsRow],
    exceptions: Sequence[ExceptionRecord],
    source_audit: Sequence[SourceAuditRecord],
    summary: dict,
    *,
    collapsed: bool = False,
) -> Path:
    template_path = Path(template_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)

    if collapsed:
        inventory = _write_collapsed_sheets(
            wb, premium_rows, claims_rows, outstanding_rows
        )
        layout = "collapsed"
    else:
        inventory = _write_class_split_sheets(
            wb,
            premium_rows,
            claims_rows,
            outstanding_rows,
            year=summary.get("year"),
            quarter=summary.get("quarter"),
        )
        layout = "class-split"

    # Update caller's summary in place when possible
    try:
        summary["output_layout"] = layout
        summary["sheet_inventory"] = inventory
    except TypeError:
        summary = dict(summary)
        summary["output_layout"] = layout
        summary["sheet_inventory"] = inventory

    write_summary_sheet(wb, summary)
    write_source_audit_sheet(wb, source_audit)
    write_exceptions_sheet(wb, exceptions)

    # Bisola-like order: SUMMARY, class sheets (already created in order),
    # then SOURCE AUDIT, EXCEPTIONS
    preferred = ["SUMMARY"] + list(inventory.keys()) + ["SOURCE AUDIT", "EXCEPTIONS"]
    _reorder_sheets(wb, preferred)

    wb.save(output_path)
    wb.close()
    return output_path
