"""Delivery of a multi-file batch (IMPL-20260929-06).

* ``separate`` (default): the per-group cleaned workbooks as written — these
  are the upload-ready files.
* ``zip``: exactly those workbooks in one archive.
* ``merged``: ONE review / convenience workbook that concatenates the
  per-group workbooks sheet by sheet in the same layout (title rows, header
  row, column order and cell formats of the cleaned workbooks). Each group's
  rows are marked in the ``MERGED REVIEW`` sheet (source period, files and the
  row range on every sheet). It is labelled NOT FOR UPLOAD; amounts are never
  summed across currencies (a non-NGN workbook's sheets keep a ``(CCY)``
  suffix).
"""
from __future__ import annotations

import re
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DELIVERY_OPTIONS = ("separate", "zip", "merged")
MERGED_SHEET = "MERGED REVIEW"
MERGED_BANNER = "REVIEW / CONVENIENCE OUTPUT — NOT FOR UPLOAD"
_SKIP_SHEETS = {"SUMMARY", "EXCEPTIONS", "SOURCE AUDIT", "SOURCE_AUDIT"}
_SERIAL_HEADERS = {"S/NO", "S/N", "SNO", "SN"}
_PERIOD_RE = re.compile(r"\bQ[1-4]\s+(?:19|20)\d{2}\b")


def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().upper().rstrip(".")


def header_row_index(ws, max_scan: int = 12) -> Optional[int]:
    """1-based header row of a cleaned sheet (the row with the S/NO cell)."""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(r, c).value) in _SERIAL_HEADERS:
                return r
    return None


def data_rows(ws) -> List[int]:
    """Row numbers of data rows (below the header, any non-empty cell)."""
    h = header_row_index(ws)
    if h is None:
        return []
    out = []
    for r in range(h + 1, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            out.append(r)
    return out


def data_sheet_counts(path: Path) -> Dict[str, int]:
    wb = load_workbook(path)
    try:
        return {ws.title: len(data_rows(ws)) for ws in wb.worksheets
                if ws.title.upper() not in _SKIP_SHEETS and ws.title != MERGED_SHEET
                and header_row_index(ws) is not None}
    finally:
        wb.close()


def upload_files(outputs: Sequence[Dict[str, Any]]) -> List[Path]:
    """The per-group upload-ready workbooks (existing files, in order, unique)."""
    seen, out = set(), []
    for o in outputs:
        p = Path(str(o.get("output_path") or ""))
        if p.is_file() and str(p) not in seen:
            seen.add(str(p))
            out.append(p)
    return out


def build_zip(outputs: Sequence[Dict[str, Any]], dest: Path) -> Path:
    """Zip exactly the per-group cleaned workbooks (flat, by file name)."""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    files = upload_files(outputs)
    names = [p.name for p in files]
    if len(set(names)) != len(names):
        raise ValueError(f"duplicate workbook names in batch: {names}")
    with zipfile.ZipFile(dest, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for p in files:
            zf.write(p, arcname=p.name)
    return dest


def _copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def _copy_sheet_frame(src_ws, dst_ws, header_row: int) -> None:
    """Title rows + header row, merged cells within them, widths, freeze."""
    for r in range(1, header_row + 1):
        for c in range(1, src_ws.max_column + 1):
            _copy_cell(src_ws.cell(r, c), dst_ws.cell(r, c))
        if src_ws.row_dimensions[r].height:
            dst_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
    for rng in src_ws.merged_cells.ranges:
        if rng.max_row <= header_row:
            dst_ws.merge_cells(str(rng))
    for key, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[key].width = dim.width
    if src_ws.freeze_panes:
        dst_ws.freeze_panes = src_ws.freeze_panes


def _sheet_title(title: str, ccy: Optional[str]) -> str:
    if ccy and str(ccy).upper() not in ("", "NGN"):
        suffix = f" ({str(ccy).upper()})"
        return (title[: 31 - len(suffix)] + suffix)[:31]
    return title[:31]


def build_merged_workbook(
    outputs: Sequence[Dict[str, Any]],
    dest: Path,
    *,
    title: str = "",
) -> Tuple[Path, Dict[str, Any]]:
    """Concatenate per-group workbooks into one review workbook.

    ``outputs`` are batch deliverable entries (``output_path``, ``group``,
    ``year``, ``quarter``, ``currency``, ``source_files``). Returns
    (path, report) where report has per-sheet row totals and the group map.
    """
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    review = wb.active
    review.title = MERGED_SHEET
    targets: Dict[str, Dict[str, Any]] = {}   # merged sheet → {ws, headers, next}
    group_map: List[List[Any]] = []
    periods: List[str] = []
    for o in outputs:
        p = Path(str(o.get("output_path") or ""))
        if not p.is_file():
            continue
        period = f"{o.get('year')} Q{o.get('quarter')}"
        if period not in periods:
            periods.append(period)
        src = load_workbook(p)
        try:
            for ws in src.worksheets:
                if ws.title.upper() in _SKIP_SHEETS:
                    continue
                h = header_row_index(ws)
                if h is None:
                    continue
                name = _sheet_title(ws.title, o.get("currency"))
                src_headers = {c: _norm(ws.cell(h, c).value) for c in range(1, ws.max_column + 1)}
                if name not in targets:
                    dst = wb.create_sheet(name)
                    _copy_sheet_frame(ws, dst, h)
                    headers = {}
                    for c, v in src_headers.items():
                        if v and v not in headers:
                            headers[v] = c
                    targets[name] = {"ws": dst, "headers": headers, "next": h + 1,
                                     "header_row": h, "serial_col": next(
                                         (c for v, c in headers.items() if v in _SERIAL_HEADERS), None)}
                t = targets[name]
                rows = data_rows(ws)
                start = t["next"]
                unmapped = sorted({v for v in src_headers.values() if v and v not in t["headers"]})
                for r in rows:
                    for c, hv in src_headers.items():
                        if not hv or hv not in t["headers"]:
                            continue
                        _copy_cell(ws.cell(r, c), t["ws"].cell(t["next"], t["headers"][hv]))
                    if t["serial_col"]:
                        t["ws"].cell(t["next"], t["serial_col"]).value = t["next"] - t["header_row"]
                    t["next"] += 1
                end = t["next"] - 1
                group_map.append([
                    o.get("group") or "", period, o.get("currency") or "", name,
                    len(rows), f"{start}–{end}" if rows else "—",
                    ", ".join(o.get("source_files") or []), p.name,
                    ", ".join(unmapped) if unmapped else "",
                ])
        finally:
            src.close()

    # Retitle the copied title rows: "… — Q4 2025" → "… — MERGED REVIEW: 2024 Q2 + 2024 Q4".
    label = "MERGED REVIEW: " + " + ".join(periods) + " (NOT FOR UPLOAD)"
    for t in targets.values():
        ws = t["ws"]
        for r in range(1, t["header_row"]):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and _PERIOD_RE.search(v):
                    ws.cell(r, c).value = _PERIOD_RE.sub(label, v, count=1)

    red = Font(bold=True, color="C00000", size=14)
    review["A1"] = MERGED_BANNER
    review["A1"].font = red
    review["A2"] = (title or "Multi-file batch") + f" — generated {datetime.now():%Y-%m-%d %H:%M}"
    review["A3"] = ("Upload the per-quarter cleaned workbooks, not this file. Rows below come "
                    "from each group's cleaned workbook unchanged (S/NO renumbered); amounts "
                    "are not summed across currencies.")
    hdr = ["GROUP", "SOURCE PERIOD", "CURRENCY", "SHEET", "ROWS", "ROW RANGE",
           "SOURCE FILES", "GROUP WORKBOOK", "COLUMNS NOT IN MERGED LAYOUT"]
    for j, v in enumerate(hdr, 1):
        cell = review.cell(5, j, v)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, row in enumerate(group_map, 6):
        for j, v in enumerate(row, 1):
            review.cell(i, j, v)
    for j, w in enumerate((34, 14, 10, 30, 8, 12, 60, 46, 30), 1):
        review.column_dimensions[get_column_letter(j)].width = w
    wb.properties.title = MERGED_BANNER
    wb.properties.subject = "Review only — per-quarter workbooks are the upload files"
    wb.save(dest)
    totals = {n: t["next"] - t["header_row"] - 1 for n, t in targets.items()}
    return dest, {"sheet_rows": totals, "group_map": group_map, "periods": periods}


def deliver(
    outputs: Sequence[Dict[str, Any]],
    option: str,
    dest_dir: Path,
    *,
    stem: str = "BATCH",
) -> List[Path]:
    """Materialise a delivery option; returns the file(s) to hand over."""
    option = (option or "separate").strip().lower()
    if option not in DELIVERY_OPTIONS:
        raise ValueError(f"delivery option must be one of {DELIVERY_OPTIONS}, got {option!r}")
    dest_dir = Path(dest_dir)
    if option == "separate":
        return upload_files(outputs)
    if option == "zip":
        return [build_zip(outputs, dest_dir / f"{stem}_cleaned_workbooks.zip")]
    path, _ = build_merged_workbook(
        outputs, dest_dir / f"{stem}_MERGED_REVIEW_NOT_FOR_UPLOAD.xlsx", title=stem)
    return [path]
