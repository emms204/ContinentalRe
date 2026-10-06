"""IMPL-20260929-08: period from row dates + de-duplicated, non-stale banners.

Fixtures are derived from a 2021 quarterly claims pack whose file names carry a
quarter but no year ('Paid  n Outstanding Claims Bord Q2..xls', 'Paid Claims
Bord Q3 .xls', 'Paid Claims Bord Q4..xls', 'Paid Claims Bord. Q1.xlsx'): the
surplus tabs repeat the stale boilerplate '1st Quarter 2015', STATEMENT says
'1ST QUARTER: 2015', LOCAL says 2014, only STATEMENT UY carries the live
'N QUARTER: 2021', and the loss dates are spread over 2016-2021 with no
majority year. No year / quarter picker: the period must be inferred.
"""
from __future__ import annotations

import random
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from cre_cleaner.core.batch import plan_batch
from cre_cleaner.core.period_infer import infer_period

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = REPO / "templates" / "TEMPLATE.xlsx"
REAL_DIR = REPO / "data" / "raw" / "newdata" / "AIICO" / "SCIB" / "2021"
REAL_NAMES = {
    1: "Paid Claims Bord. Q1.xlsx",
    2: "Paid  n Outstanding Claims Bord Q2..xls",
    3: "Paid Claims Bord Q3 .xls",
    4: "Paid Claims Bord Q4..xls",
}
ORD = {1: "1ST", 2: "2ND", 3: "3RD", 4: "4TH"}
# Loss-date year spread seen in the real pack (no year above 50 %).
LOSS_SPREAD = {2021: 13, 2020: 25, 2019: 13, 2018: 9, 2017: 4, 2016: 2}


def _claims_rows(title, seed):
    rng = random.Random(seed)
    rows = [["COMPANY"], [title], ["S/NO", "DATE OF LOSS", "MONTH", "NAME OF INSURED", "CLAIM PAID"]]
    i = 0
    for y, n in LOSS_SPREAD.items():
        for _ in range(n):
            i += 1
            rows.append([i, datetime(y, rng.randint(1, 12), rng.randint(1, 28)), "JANUARY",
                         f"INSURED {i}", 1000 + i])
    return rows


def _pack(path: Path, q: int, *, live_uy=True, stale_label_quarter=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "LOCAL"
    for r in (["COMPANY"], ["ACCOUNTING QUARTER: 1ST", "1ST QUARTER:", 2014.0],
              ["CLASS", "PREMIUM"], ["FIRE", 10]):
        ws.append(r)
    ws = wb.create_sheet("STATEMENT")
    for r in (["COMPANY"], ["ACCOUNTING QUARTER:", "1ST QUARTER:", 2015.0],
              ["CLASS", "PREMIUM"], ["FIRE", 10]):
        ws.append(r)
    if live_uy:
        ws = wb.create_sheet("STATEMENT UY")
        for r in (["COMPANY"], ["ACCOUNTING QUARTER:", f"{ORD[q]} QUARTER:", 2021.0],
                  ["CLASS", "PREMIUM"], ["FIRE", 10]):
            ws.append(r)
    for cls in ("FIRE", "MARINE CARGO", "ENGINEERING"):
        ws = wb.create_sheet(f"{cls} SURPLUS")
        for r in (["COMPANY"], ["1st Quarter 2015"], ["CLASS", "PREMIUM"], [cls, 10]):
            ws.append(r)
    if stale_label_quarter:  # year right, quarter left at Q1 on every file
        ws = wb.create_sheet("FORMAT")
        for r in (["COMPANY"], ["1st QUARTER 2021"], ["CLASS", "PREMIUM"], ["FIRE", 10]):
            ws.append(r)
    for k, tab in enumerate(("Fire Paid Claim", "Engineering", "Marine")):
        ws = wb.create_sheet(tab)
        for r in _claims_rows(f"{tab.upper()} PAID CLAIMS", seed=q * 10 + k):
            ws.append(r)
    wb.save(path)
    return path


def _fixture_names():
    # Same stems (spacing kept) as the real uploads; saved as .xlsx.
    return {q: Path(n).stem + ".xlsx" for q, n in REAL_NAMES.items()}


@pytest.mark.parametrize("q", [1, 2, 3, 4])
def test_filename_quarter_plus_row_year_beats_stale_banners(tmp_path, q):
    p = _pack(tmp_path / _fixture_names()[q], q, live_uy=(q != 4))
    r = infer_period(p)
    assert (r.year, r.quarter) == (2021, q), r.evidence[:4]
    assert r.ok and not r.filename_override
    reasons = {(s, c) for s, c, _d in r.notes}
    assert ("INFO", "period_banner_stale") in reasons       # 2014 / 2015 banners dropped
    assert ("WARN", "period_banner_conflict") not in reasons


def test_four_quarter_uploads_become_four_jobs(tmp_path):
    up = tmp_path / "upload"
    up.mkdir()
    files = [_pack(up / n, q, live_uy=(q != 4)) for q, n in _fixture_names().items()]
    plan = plan_batch(files, cedant="AIICO", broker="SCIB")
    got = sorted((g.year, g.quarter, len(g.files)) for g in plan.groups)
    assert got == [(2021, 1, 1), (2021, 2, 1), (2021, 3, 1), (2021, 4, 1)]


def test_identical_banner_text_counts_once(tmp_path):
    """Three tabs repeating '1st Quarter 2015' are one vote, not three."""
    wb = Workbook()
    ws = wb.active
    ws.title = "A SURPLUS"
    ws.append(["1st Quarter 2015"])
    for t in ("B SURPLUS", "C SURPLUS"):
        wb.create_sheet(t).append(["1st Quarter 2015"])
    wb.create_sheet("UY").append(["ACCOUNTING QUARTER: 2ND QUARTER: 2021"])
    wb.create_sheet("SUMMARY").append(["2ND QUARTER 2021 SUMMARY"])
    p = tmp_path / "CLAIMS BORD.xlsx"
    wb.save(p)
    r = infer_period(p)
    assert (r.year, r.quarter) == (2021, 2)
    # banners disagree on the year (no row dates to mark 2015 stale) → WARN
    assert any(c == "period_banner_conflict" for s, c, _d in r.notes if s == "WARN")


def test_reporting_date_rows_beat_banner_with_conflict_warn(tmp_path):
    """C4: banner 4th Quarter 2020 vs payment dates mostly in 2021 Q1, no file
    name quarter → row year wins, WARN period_banner_conflict (never silent)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CLAIMS"
    ws.append(["4TH QUARTER 2020 CLAIMS BORDEREAU"])
    ws.append(["S/NO", "PAYMENT DATE", "INSURED", "CLAIM PAID"])
    for i in range(12):
        ws.append([i + 1, datetime(2021, 1 + i % 3, 5), f"I{i}", 100])
    p = tmp_path / "CLAIMS BORDEREAU.xlsx"
    wb.save(p)
    r = infer_period(p)
    assert (r.year, r.quarter) == (2021, 1)
    warn = [d for s, c, d in r.notes if (s, c) == ("WARN", "period_banner_conflict")]
    assert warn and "2020" in warn[0]
    assert r.confidence == "medium"


def test_banner_tie_without_filename_quarter_is_period_ambiguous(tmp_path):
    """C4: two live banners disagree, row dates settle nothing, file name names
    no period → WARN period_banner_conflict + ERROR period_ambiguous."""
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "raw"
    raw.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "STATEMENT"
    ws.append(["2ND QUARTER 2020"])
    ws.append(["S/NO", "DATE OF LOSS", "INSURED", "CLAIM PAID"])
    for i, y in enumerate([2019, 2020, 2021, 2018, 2020, 2021]):
        ws.append([i + 1, datetime(y, 3, 3), f"I{i}", 100])
    wb.create_sheet("UY").append(["3RD QUARTER 2021"])
    wb.save(raw / "CLAIMS BORDEREAU.xlsx")
    r = infer_period(raw)
    assert not r.ok and r.year is None
    assert any((s, c) == ("WARN", "period_banner_conflict") for s, c, _d in r.notes)
    res = run_pipeline(cedant="UNITRUST", broker="AGRIC", raw_dir=raw, template=TEMPLATE,
                       out_dir=tmp_path / "out", base_dir=REPO)
    reasons = {(e.severity, e.reason) for e in res.exceptions}
    assert ("WARN", "period_banner_conflict") in reasons
    assert ("ERROR", "period_ambiguous") in reasons
    assert not list((tmp_path / "out").glob("*_cleaned.xlsx"))


def test_stale_banner_is_logged_in_pipeline(tmp_path):
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "raw"
    raw.mkdir()
    _pack(raw / _fixture_names()[3], 3)
    res = run_pipeline(cedant="UNITRUST", broker="AGRIC", raw_dir=raw, template=TEMPLATE,
                       out_dir=tmp_path / "out", base_dir=REPO)
    reasons = {(e.severity, e.reason) for e in res.exceptions}
    assert ("INFO", "period_banner_stale") in reasons
    assert (res.summary.get("year"), res.summary.get("quarter")) == (2021, 3)


@pytest.mark.skipif(not all((REAL_DIR / n).exists() for n in REAL_NAMES.values()),
                    reason="2021 claims pack not present")
def test_real_2021_quarter_uploads_plan_four_jobs(tmp_path):
    import shutil
    up = tmp_path / "upload"  # no year folder: the upload keeps file names only
    up.mkdir()
    files = []
    for n in REAL_NAMES.values():
        shutil.copy2(REAL_DIR / n, up / n)
        files.append(up / n)
    plan = plan_batch(files, cedant="AIICO", broker="SCIB")
    got = sorted((g.year, g.quarter, len(g.files)) for g in plan.groups)
    assert got == [(2021, 1, 1), (2021, 2, 1), (2021, 3, 1), (2021, 4, 1)]
