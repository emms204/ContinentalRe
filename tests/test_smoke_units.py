"""Minimal unit tests: header detect, negative preserve, monthly merge order."""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from cre_cleaner.core.detect import find_header_row, detect_sheet_type
from cre_cleaner.core.normalize import parse_number, as_text_id, parse_period
from cre_cleaner.config import QUARTER_MONTHS


def test_find_header_row_premium():
    rows = [
        [None, None, "OWN RETENTION", None, None, "1SURP"],
        [None, "INSURED", "POLICY NO", "PERIOD OF INSURANCE", "SUM INSURED", "PREMIUM", "PPN", "SUM INSURED", "Premium"],
        [None, "ACME", "P/1", "01/01/2025 - 31/12/2025", 1000, 50, 10, 100, 5],
    ]
    assert find_header_row(rows, "premium") == 1


def test_detect_paid_vs_outstanding():
    from cre_cleaner.adapters.aiico_ark import AiicoArkAdapter
    aiico = AiicoArkAdapter().sheet_type_rules  # IMPL-20260929-04: AIICO tab rules
    assert detect_sheet_type("MISC. ACC PAID JANUARY", [], aiico) == "paid"
    assert detect_sheet_type("MARINE HULL OUTSTANDING CLAIM", []) == "outstanding"
    assert detect_sheet_type("FIRE 2ND", [[None, "INSURED", "POLICY NO", "PREMIUM"]]) == "premium"
    # REPORT risk 2: bare OS / O/S must not be typed as paid (AIICO SCIB tabs)
    assert detect_sheet_type("marine os claims", [], aiico) == "outstanding"
    assert detect_sheet_type("Fire O/S Claims", [], aiico) == "outstanding"


def test_sheet_type_rules_are_adapter_owned():
    """IMPL-20260929-04: AIICO / HEIRS JOMOLA tab rules are not generic."""
    from cre_cleaner.adapters import get_adapter
    assert detect_sheet_type("MISC. ACC PAID JANUARY", []) != "paid"
    assert detect_sheet_type("FIRE PRODUCTION", []) == "premium"
    assert detect_sheet_type("FIRE PRODUCTION", [],
                             get_adapter("HEIRS", "JOMOLA").sheet_type_rules) == "skip"


def test_quarters_in_text_does_not_eat_year_digit():
    from cre_cleaner.adapters.base import quarters_in_text
    # REPORT risk 1: "Qtr 2021" must not add quarter 2 from the year
    assert quarters_in_text("4TH Qtr 2021 Claims Bord.xls") == {4}
    assert quarters_in_text("1ST QTR 2020 CLAIMS.xls") == {1}
    assert quarters_in_text("3RD QTR 2024 PREMIUM.xlsx") == {3}
    assert quarters_in_text("QTR1 2023") == {1}


def test_parse_number_trailing_minus_and_currency():
    assert parse_number("1,234-") == -1234.0
    assert parse_number("₦1,234,567.89") == 1234567.89
    assert parse_number("NGN 1,000.00") == 1000.0
    assert parse_number("(1,234.50)") == -1234.5


def test_claims_share_headers_not_mapped_to_amounts():
    from cre_cleaner.core.map_columns import map_simple_columns, CLAIMS_ALIASES, merge_group_subheaders
    # HEIRS: TREATY PPN% must not become TREATY AMOUNT
    headers = [
        "CLAIM NO", "POLICY NO", "NAME OF INSURED CLAIMANT",
        "TOTAL AMOUNT PAID", "TREATY PPN%", "TREATY RECOVERY",
    ]
    cm = map_simple_columns(headers, CLAIMS_ALIASES)
    assert cm.get("total_claims") == 3
    assert cm.get("amount_treaty") == 5
    # LASACO-style band + %/AMOUNT subrow
    merged = merge_group_subheaders(
        ["GROSS LOSS RESERVE", "RETENTION", None, "TREATY", None],
        [None, "%", "AMOUNT", "%", "AMOUNT"],
    )
    assert merged[1:5] == ["RETENTION PPN", "RETENTION AMOUNT", "TREATY PPN", "TREATY AMOUNT"]
    cm2 = map_simple_columns(merged, CLAIMS_ALIASES)
    assert cm2.get("amount_ret") == 2
    assert cm2.get("amount_treaty") == 4


def test_unitrust_ppn_after_si_prem_blocks():
    from cre_cleaner.adapters import get_adapter
    header = [
        "INSURED NAME", "POLICY NO", "OUR SHARE SUM INSURED", "OUR SHARE GROSS PREMIUM",
        "RETENTION SUM INSURED", "RETENTION GROSS PREMIUM", "RETENTION PPN",
        "TREATY SUM INSURED", "TREATY GROSS PREMIUM", "TREATY PPN",
    ]
    cm = get_adapter("UNITRUST", "AGRIC").detect_premium_allocation_blocks(header, None)
    # IMPL-20260929-04 (Emmanuel): our-share columns are never TSI / GP; with no
    # 100% column the gross fields stay blank (never another band's figures).
    assert cm.get("gross_premium") is None and cm.get("sum_insured") is None, cm.mapping
    assert cm.ret_ppn == 6 and cm.ret_si == 4 and cm.ret_prem == 5
    assert cm.sur_ppn == 9 and cm.sur_si == 7 and cm.sur_prem == 8


def test_unitrust_agric_bisola_gold_columns():
    """Unitrust Agric: policy SI/GP + NET RETENTION/TREATY QUOTA SHARE blank bands.

    Bisola gold Q2 2024: TOTAL SUM INSURED = Sum Insured (not Our share SI),
    GROSS PREMIUM = Gross Premium (not Our Share of G. Prem.), RET/TREATY =
    % then SI then premium in the next two blank-header columns.
    """
    from cre_cleaner.adapters import get_adapter
    detect_premium_allocation_blocks = get_adapter("UNITRUST", "AGRIC").detect_premium_allocation_blocks
    header = [
        "Policy No ", "Insured ", "Policy Class ", "Risktype ", "Branch ", "Debit Note ",
        "Company's Share ", "MPL % ", "Insurance Period ", "Sum Insured ", "Our share SI ",
        "Gross Premium ", None, "Our Share of G. Prem. ", "NET RETENTION ", None, None,
        "TREATY QUOTA SHARE", None, None,
    ]
    cm = detect_premium_allocation_blocks(header, None)
    assert cm.get("sum_insured") == 9, cm.mapping  # Sum Insured, not Our share SI
    assert cm.get("gross_premium") == 11, cm.mapping  # Gross Premium, not Our Share of G. Prem.
    assert cm.ret_ppn == 14 and cm.ret_si == 15 and cm.ret_prem == 16
    assert cm.sur_ppn == 17 and cm.sur_si == 18 and cm.sur_prem == 19


def test_duplicate_copy_sheet_detection():
    from cre_cleaner.core.quarterly import _is_duplicate_copy_sheet
    names = ["CLAIMS RECOVERY", "CLAIMS RECOVERY (2)", "Sheet1"]
    assert _is_duplicate_copy_sheet("CLAIMS RECOVERY (2)", names)
    assert not _is_duplicate_copy_sheet("CLAIMS RECOVERY", names)
    assert not _is_duplicate_copy_sheet("CLAIMS RECOVERY (2)", ["CLAIMS RECOVERY (2)"])


def test_flag_duplicate_premium_exact():
    from cre_cleaner.models import PremiumRow
    from cre_cleaner.core.reconcile import flag_duplicate_premium
    from datetime import datetime
    a = PremiumRow(policy_no="P1", name_of_insured="A",
                   period_from=datetime(2024, 1, 1), period_to=datetime(2024, 12, 31),
                   gross_premium=100, total_sum_insured=1000)
    b = PremiumRow(policy_no="P1", name_of_insured="A",
                   period_from=datetime(2024, 1, 1), period_to=datetime(2024, 12, 31),
                   gross_premium=100, total_sum_insured=1000)
    c = PremiumRow(policy_no="P1", name_of_insured="A",
                   period_from=datetime(2024, 1, 1), period_to=datetime(2024, 6, 30),
                   gross_premium=100, total_sum_insured=1000)
    # Row-level repeats are kept as in the source and are not flagged.
    assert flag_duplicate_premium([a, b]) == []
    assert flag_duplicate_premium([a, c]) == []


def test_aiico_rtntn_labelled_blocks():
    from cre_cleaner.adapters.aiico_ark import AiicoArkAdapter
    detect_premium_allocation_blocks = AiicoArkAdapter().detect_premium_allocation_blocks
    header = [
        "INSURED", "POLICY NO", "COVER START", "COVER END", "SUM INSURED", "PREMIUM",
        "RTNTN PPN", "RTNTN SUM INSURED", "RTNTN PREMIUM",
        "QUOTA SHARE PPN", "QUOTA SHARE SUM INSURED", "QUOTA SHARE PREMIUM",
    ]
    cm = detect_premium_allocation_blocks(header, None)
    assert cm.get("period_from") == 2 and cm.get("period_to") == 3
    assert cm.ret_ppn == 6 and cm.ret_si == 7 and cm.ret_prem == 8
    assert cm.sur_ppn == 9 and cm.sur_si == 10 and cm.sur_prem == 11


def test_preserve_negatives():
    assert parse_number(-10846229.6) == -10846229.6
    assert parse_number("(1,234.50)") == -1234.50
    assert parse_number("NIL") is None


def test_text_ids():
    assert as_text_id("503/00038/19/TQ/HO") == "503/00038/19/TQ/HO"
    assert as_text_id(2025.0) == "2025"


def test_period_split():
    a, b = parse_period("05/10/2024 - 04/04/2025")
    assert a is not None and b is not None
    assert a.day == 5 and a.month == 10 and a.year == 2024
    assert b.day == 4 and b.month == 4 and b.year == 2025


def test_parse_period_single_date_not_copied_to_both():
    from datetime import datetime
    a, b = parse_period(datetime(2024, 6, 15))
    assert a is not None and a.day == 15 and a.month == 6
    assert b is None
    a, b = parse_period("15/06/2024")
    assert a is not None and b is None


def test_unlabelled_allocation_first_retention_second_treaty():
    from cre_cleaner.core.map_columns import detect_premium_allocation_blocks
    # No group labels; two PPN/SI/Prem triples → first RETENTION, second TREATY
    header = [
        "INSURED", "POLICY",
        "PPN", "SUM INSURED", "PREMIUM",
        "PPN", "SUM INSURED", "PREMIUM",
    ]
    cm = detect_premium_allocation_blocks(header, group_row=None)
    assert cm.ret_ppn == 2 and cm.sur_ppn == 5
    assert cm.ret_si == 3 and cm.sur_si == 6
    assert cm.ret_prem == 4 and cm.sur_prem == 7
    assert cm.sur_label == "TREATY"


def test_single_unlabelled_block_is_treaty_not_retention():
    from cre_cleaner.core.map_columns import detect_premium_allocation_blocks
    header = ["INSURED", "PPN", "SUM INSURED", "PREMIUM"]
    cm = detect_premium_allocation_blocks(header, group_row=None)
    assert cm.ret_ppn is None and cm.ret_si is None and cm.ret_prem is None
    assert cm.sur_ppn == 1 and cm.sur_si == 2 and cm.sur_prem == 3
    assert cm.sur_label == "TREATY"


def test_duplicate_claims_differ_by_period_not_flagged():
    from datetime import datetime
    from cre_cleaner.models import ClaimsRow
    from cre_cleaner.core.reconcile import flag_duplicate_claims
    rows = [
        ClaimsRow(claim_no="C1", policy_no="P1", date_of_loss=datetime(2024, 1, 1),
                  period_from=datetime(2024, 1, 1), period_to=datetime(2024, 3, 31),
                  total_claims=100),
        ClaimsRow(claim_no="C1", policy_no="P1", date_of_loss=datetime(2024, 1, 1),
                  period_from=datetime(2024, 4, 1), period_to=datetime(2024, 6, 30),
                  total_claims=100),
    ]
    assert flag_duplicate_claims(rows, "CLAIMS") == []
    same = [
        rows[0],
        ClaimsRow(claim_no="C1", policy_no="P1", date_of_loss=datetime(2024, 1, 1),
                  period_from=datetime(2024, 1, 1), period_to=datetime(2024, 3, 31),
                  total_claims=100),
    ]
    # Exact claim repeats are also left unflagged (passthrough).
    assert flag_duplicate_claims(same, "CLAIMS") == []


def test_blank_from_to_flagged():
    from cre_cleaner.models import ClaimsRow, AuditMeta
    from cre_cleaner.core.reconcile import check_row_dates
    row = ClaimsRow(
        claim_no="C1", policy_no="P1", insured="A",
        date_of_loss=__import__("datetime").datetime(2024, 2, 1),
        period_from=None, period_to=None, total_claims=10,
        audit=AuditMeta(source_filename="x.xlsx", source_sheet="S", source_row=2),
    )
    exc, _counts = check_row_dates([], [row], [], year=2024, quarter=1)
    reasons = {e.reason for e in exc}
    assert "period_from_missing" in reasons
    assert "period_to_missing" in reasons


def test_discovery_qtr1_bord_jan_and_monthly_claims():
    import tempfile
    from cre_cleaner.adapters.base import quarters_in_text
    from cre_cleaner.adapters.aiico_ark import AiicoArkAdapter
    assert 1 in quarters_in_text("QTR1")
    assert 1 in quarters_in_text("QTR1 RETURNS.xlsx")
    ad = AiicoArkAdapter()
    month, why = ad.classify_premium_candidate("BORD JAN 2024.xlsx", 2024)
    assert month == 1 and why == "ok"
    month, why = ad.classify_premium_candidate("QTR1.xlsx", 2024)
    assert month is None and "quarterly" in why
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "QTR1.xlsx").write_bytes(b"PK")  # name-only discovery
        (root / "BORD JAN 2024.xlsx").write_bytes(b"PK")
        (root / "JANUARY CLAIMS 2024.xls").write_bytes(b"PK")
        q = ad.discover_quarterly_premium_files(root, 2024, 1)
        assert any(p.name == "QTR1.xlsx" for p in q)
        months = ad.discover_premium_files(root, 2024, 1)
        assert any(m == 1 and p.name.startswith("BORD") for m, p in months)
        claims = ad.discover_claims_files(root, 2024, 1)
        assert any("JANUARY CLAIMS" in p.name for p in claims)


def test_api_safe_filename_and_labels():
    from fastapi import HTTPException
    from cre_cleaner.api import safe_upload_basename, sanitize_label
    assert safe_upload_basename("../../x.xlsx") == "x.xlsx"
    assert safe_upload_basename("foo/bar/BORD JAN.xlsx") == "BORD_JAN.xlsx"
    try:
        safe_upload_basename("Q2 2024 PREMIUM.pdf")
        assert False, "PDF paused in Phase 1"
    except HTTPException as e:
        assert e.status_code == 400
    try:
        safe_upload_basename("../../etc/passwd")
        assert False, "expected reject"
    except HTTPException as e:
        assert e.status_code == 400
    assert sanitize_label("aiico", "cedant") == "AIICO"
    try:
        sanitize_label("../evil", "cedant")
        assert False, "expected reject"
    except HTTPException as e:
        assert e.status_code == 400
    try:
        sanitize_label("a/b", "broker")
        assert False, "expected reject"
    except HTTPException as e:
        assert e.status_code == 400


def test_bordereau_type_and_field_completeness():
    from cre_cleaner.pipeline import normalize_bordereau_type, field_completeness
    from cre_cleaner.models import PremiumRow
    assert normalize_bordereau_type("Premium") == "premium"
    assert normalize_bordereau_type("ost") == "outstanding"
    assert normalize_bordereau_type(None) == "all"
    rows = [
        PremiumRow(policy_no="P1", gross_premium=100, ret_prem=40, sur_ppn=60),
        PremiumRow(policy_no="", gross_premium=None, ret_prem=None, sur_prem=None),
    ]
    fc = field_completeness(rows, [], [])
    assert fc["premium_rows"] == 2
    assert fc["premium_policy_no_pct"] == 50.0
    assert fc["premium_gross_pct"] == 50.0
    assert fc["premium_retention_pct"] == 50.0
    assert fc["premium_treaty_pct"] == 50.0


def test_quarter_months_order():
    assert QUARTER_MONTHS[1] == [1, 2, 3]


def test_class_label_mapping():
    from cre_cleaner.adapters.aiico_ark import AiicoArkAdapter
    from cre_cleaner.core.class_labels import normalize_class_label, class_sheet_title
    aiico = AiicoArkAdapter().class_map()  # IMPL-20260929-04: AIICO subclasses
    assert normalize_class_label("FIRE") == "Fire"
    assert normalize_class_label("GOODS IN TRANSIT", aiico) == "General Accident"
    assert normalize_class_label("ALL RISKS", aiico) == "General Accident"
    assert normalize_class_label("GOODS IN TRANSIT") == "Goods In Transit"
    assert class_sheet_title("PUBLIC/PRODUCT LIABILITY", "CLAIMS") == "PUBLIC-PRODUCT LIABILITY - CLAIMS"
    assert normalize_class_label("MARINE HULL") == "Marine Hull"
    assert class_sheet_title("Fire", "CLAIMS") == "Fire - CLAIMS"


def test_banner_class_and_typos():
    from cre_cleaner.adapters.aiico_ark import AiicoArkAdapter
    from cre_cleaner.core.class_labels import banner_class_label, is_unresolved_class
    from cre_cleaner.core.class_labels import normalize_class_label as _n
    normalize_class_label = lambda raw: _n(raw, AiicoArkAdapter().class_map())  # noqa: E731
    assert banner_class_label("FIRE PAID CLAIM") == "Fire"
    assert banner_class_label("ENGINERRING") == "Engineering"
    assert banner_class_label("ENGINEERING PAID CLAIM") == "Engineering"
    # Bare MARINE is too broad for Hull vs Cargo — banner still matches, rows go to exceptions.
    assert banner_class_label("MARINE PAID CLAIM") == "MARINE"
    assert is_unresolved_class("MARINE")
    assert banner_class_label("MARINE CARGO") == "Marine Cargo"
    assert banner_class_label("MOTOR") == "Motor"
    assert banner_class_label("2ND SURP PAID CLAIM") == ""
    assert banner_class_label("558047.14") == ""
    assert normalize_class_label("GEN ACCIENT") == "General Accident"
    assert normalize_class_label("GEN. ACCIDENT") == "General Accident"
    assert normalize_class_label("GEN ACC") == "General Accident"
    assert normalize_class_label("MOTOR") == "Motor"
    assert normalize_class_label("CASUALTY") == "General Accident"


def test_ppn_share_bisola_style():
    # IMPL-20260929-04: AIICO-only derivation (flagged 'Calculated' by the parser)
    from cre_cleaner.adapters.aiico_common import aiico_ppn_share as _ppn_share
    # UBA OST gold: 7644.71 / 32490 ≈ 0.235294...
    assert abs(_ppn_share(7644.71, 32490) - 0.2352942444) < 1e-9
    assert abs(_ppn_share(24845.29, 32490) - 0.7647057556) < 1e-9
    assert _ppn_share(None, 32490) == 0.0
    assert _ppn_share(100, 0) is None
    assert _ppn_share(100, None) is None


def test_currency_and_split_helpers():
    from cre_cleaner.core.normalize import (
        currency_codes, currency_code, currency_filter_code, prefer_named_over_fcy,
    )
    from cre_cleaner.core.reconcile import check_row_splits
    from cre_cleaner.models import PremiumRow, AuditMeta
    assert currency_codes("Q2 2025 DOLLAR BORDEREAU") == {"USD"}
    assert currency_code("Naira") == "NGN"
    assert currency_filter_code("Currency Filter ( NAIRA at 1 )") == "NGN"
    assert currency_filter_code("Currency Filter ( USD )") == "USD"
    assert currency_filter_code("Currency Filter ( None )") == ""
    assert prefer_named_over_fcy({"FCY", "NGN"}) == {"NGN"}
    assert prefer_named_over_fcy({"FCY"}) == {"FCY"}
    # Title says FOREIGN CURRENCY but filter names Naira → named wins in the set
    both = currency_codes("NAICOM BORDEREAUX REPORT - FOREIGN CURRENCY") | currency_codes(
        "Currency Filter ( NAIRA at 1 )"
    )
    assert prefer_named_over_fcy(both) == {"NGN"}
    row = PremiumRow(
        policy_no="P1", name_of_insured="A",
        gross_premium=100, ret_prem=40, sur_prem=60, fac_prem=0,
        ret_ppn=40, sur_ppn=60, fac_ppn=0,
        audit=AuditMeta(currency="NGN"),
    )
    exc, counts = check_row_splits([row], [], [])
    assert counts["premium"]["ok"] == 1
    assert not exc


def test_unitrust_foreign_title_stays_naira_with_currency_filter():
    """Unitrust Agric: FOREIGN CURRENCY title + Currency Filter (NAIRA) → one NGN book."""
    from cre_cleaner.core.quarterly import _SheetCurrency, _workbook_currency_filter_hint
    from cre_cleaner.models import ExceptionRecord
    from pathlib import Path

    premium_top = [
        ["UNITRUST INSURANCE COMPANY LIMITED"],
        ["NAICOM BORDEREAUX REPORT - FOREIGN CURRENCY"],
        ["Currency Filter ( NAIRA at 1 )"],
        ["Policy No", "Sum Insured", "Gross Premium", "NET RETENTION"],
    ]
    claims_top = [
        ["CLAIMS PAID REPORT WITH RE-INSURANCE - FOREIGN CURRENCY"],
        ["Branch", "Policy No.", "Total Claims Paid", "Net Retention"],
    ]
    sheets = {"AGRIC PREMIUM": premium_top, "AGRIC PAID CLAIM": claims_top}
    hint = _workbook_currency_filter_hint(sheets)
    assert hint == "NGN"
    exc: list = []
    prem_ccy = _SheetCurrency(Path("x.xlsx"), "AGRIC PREMIUM", premium_top, exc)
    assert prem_ccy.code == "NGN" and prem_ccy.source == "currency filter"
    claim_exc: list = []
    claim_ccy = _SheetCurrency(
        Path("x.xlsx"), "AGRIC PAID CLAIM", claims_top, claim_exc,
        file_currency_hint=hint,
    )
    assert claim_ccy.code == "NGN"
    assert claim_ccy.source == "workbook currency filter"
    assert any(e.reason == "currency_workbook_filter" for e in claim_exc)


def test_period_infer_from_filenames():
    from cre_cleaner.core.period_infer import infer_period_from_labels
    r = infer_period_from_labels(["Q2 2024 PREMIUM BORDEREAU.xlsx", "Claims Q2 2024.xls"])
    assert r.year == 2024 and r.quarter == 2 and r.ok
    r = infer_period_from_labels(["APRIL PREMIUM 2025 LOCAL.xlsx"])
    assert r.year == 2025 and r.quarter == 2


def test_period_infer_from_date_columns():
    """Year/quarter come from sheet date columns, not the filename."""
    import tempfile
    from openpyxl import Workbook
    from cre_cleaner.core.period_infer import infer_period

    wb = Workbook()
    ws = wb.active
    ws.title = "PAID"
    ws.append(["INSURED", "DATE OF LOSS", "CLAIM NO", "TOTAL CLAIMS"])
    ws.append(["A", datetime(2024, 4, 12), "C1", 100])
    ws.append(["B", datetime(2024, 5, 3), "C2", 200])
    ws.append(["C", datetime(2024, 6, 20), "C3", 50])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "mystery_bordereau.xlsx"  # no year/quarter in name
        wb.save(path)
        r = infer_period(Path(tmp))
    assert r.ok and r.source == "content"
    assert r.year == 2024 and r.quarter == 2


def test_period_infer_from_to_banner_beats_row_dates():
    """Unitrust Period Date From/To is the reporting quarter, not Start Date."""
    import tempfile
    from openpyxl import Workbook
    from cre_cleaner.core.period_infer import (
        _period_from_from_to_banner,
        infer_period,
    )

    assert _period_from_from_to_banner(
        "From:[April 01, 2023 ] To [ June 30, 2023 ]"
    ) == (2023, 2)
    assert _period_from_from_to_banner(
        "\xa0From:\xa0[April 01, 2023 ] \xa0\xa0To\xa0\xa0[ June 30, 2023 ] "
    ) == (2023, 2)

    wb = Workbook()
    ws = wb.active
    ws.title = "AGRIC PAID CLAIM"
    # Rows 1–4 empty-ish; row 5 = Period Date banner (0-index 4 in reader).
    ws.append(["UNITRUST"])
    ws.append(["CLAIMS PAID REPORT"])
    ws.append(["Print Date"])
    ws.append([])
    ws.append([
        "Period Date ( Payment Date )",
        *[""] * 28,
        "From:[April 01, 2023 ] To [ June 30, 2023 ]",
    ])
    for _ in range(6):
        ws.append([])
    ws.append([
        "Branch", "Pol. Class", "Start Date", "Date Of Loss", "Claim No",
    ])
    ws.append([
        "HEAD OFFICE", "Agriculture", datetime(2023, 2, 13),
        datetime(2023, 2, 14), "C1",
    ])
    with tempfile.TemporaryDirectory() as tmp:
        # Filename without quarter — banner alone must yield Q2.
        path = Path(tmp) / "TREATY AGRIC 2023.xlsx"
        wb.save(path)
        r = infer_period(Path(tmp))
    assert r.ok and r.year == 2023 and r.quarter == 2
    assert r.source == "content"
    assert any("banner" in e.lower() for e in r.evidence)


def test_period_infer_filename_beats_cover_date_vote():
    """2ND QTR filename wins over cover-start dates that vote Q1 → discovery finds the file."""
    import tempfile
    from openpyxl import Workbook
    from cre_cleaner.core.period_infer import infer_period

    wb = Workbook()
    ws = wb.active
    ws.title = "PAID"
    ws.append(["INSURED", "START DATE", "DATE OF LOSS", "CLAIM NO"])
    ws.append(["A", datetime(2023, 2, 13), datetime(2023, 2, 14), "C1"])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "2ND QTR TREATY AGRIC 2023.xlsx"
        wb.save(path)
        r = infer_period(Path(tmp))
    assert r.ok and r.year == 2023 and r.quarter == 2
    # IMPL-20260929-05: the file name replaces the weaker date vote and the
    # override is recorded (WARN period_filename_override in the pipeline).
    assert r.source == "filename" and r.filename_override


def test_adapter_registry_no_fallback():
    from cre_cleaner.adapters import get_adapter, list_adapters
    from cre_cleaner.adapters.base import UnsupportedCedantError
    assert get_adapter("AIICO", "ARK").verified is True
    assert get_adapter("CUSTODIAN", "SCIB").verified is False
    assert get_adapter("AXA", "DIRECT").cedant == "AXA"
    assert get_adapter("CHI 2", "SCIB").cedant == "CHI"
    assert get_adapter("UNITRUST", "UNITRUST & ARK").broker == "ARK"
    assert len(list_adapters()) >= 15
    try:
        get_adapter("UNKNOWN", "X")
        assert False, "expected UnsupportedCedantError"
    except UnsupportedCedantError:
        pass


def test_claims_leading_blank_default():
    from openpyxl import Workbook
    from cre_cleaner.io.excel import claims_col_map, write_claims_rows
    from cre_cleaner.models import ClaimsRow
    assert min(claims_col_map().values()) == 2
    assert min(claims_col_map(False).values()) == 1
    ws = Workbook().active
    write_claims_rows(ws, [ClaimsRow(insured="X", policy_no="P1", claim_no="C1", total_claims=1.0)],
                      title="T")
    assert ws["B1"].value == "T"
    assert all(c.value is None for c in ws["A"])
    assert ws.max_column == 18  # A blank + headers B..R (TEMPLATE; no SUM INSURED)
    hdr = [ws.cell(4, c).value for c in range(2, 19)]
    assert "SUM INSURED" not in hdr
    assert hdr[8] == "TO" and hdr[9] == "TOTAL CLAIMS"


def test_claims_ppn_uses_percent_number_format():
    """Claims PPN are fractions; Excel 0.00% displays them like Bisola (6.21%)."""
    from openpyxl import Workbook
    from cre_cleaner.io.excel import write_claims_rows
    from cre_cleaner.models import ClaimsRow
    ws = Workbook().active
    write_claims_rows(
        ws,
        [ClaimsRow(
            insured="X", policy_no="P1", claim_no="C1",
            total_claims=188771.19, ppn_ret=0.06214301027609138,
            amount_ret=11730.81, ppn_treaty=0.2758399732501554, amount_treaty=52070.64,
        )],
        title="T",
    )
    assert ws.cell(5, 12).value == 0.06214301027609138  # PPN RET %
    assert ws.cell(5, 12).number_format == "0.00%"
    assert ws.cell(5, 14).number_format == "0.00%"  # PPN TREATY %
    assert ws.cell(5, 16).number_format == "0.00%"  # PPN FAC %


if __name__ == "__main__":
    test_find_header_row_premium()
    test_detect_paid_vs_outstanding()
    test_sheet_type_rules_are_adapter_owned()
    test_quarters_in_text_does_not_eat_year_digit()
    test_parse_number_trailing_minus_and_currency()
    test_claims_share_headers_not_mapped_to_amounts()
    test_unitrust_ppn_after_si_prem_blocks()
    test_unitrust_agric_bisola_gold_columns()
    test_duplicate_copy_sheet_detection()
    test_flag_duplicate_premium_exact()
    test_aiico_rtntn_labelled_blocks()
    test_preserve_negatives()
    test_text_ids()
    test_period_split()
    test_parse_period_single_date_not_copied_to_both()
    test_unlabelled_allocation_first_retention_second_treaty()
    test_single_unlabelled_block_is_treaty_not_retention()
    test_duplicate_claims_differ_by_period_not_flagged()
    test_blank_from_to_flagged()
    test_discovery_qtr1_bord_jan_and_monthly_claims()
    test_api_safe_filename_and_labels()
    test_bordereau_type_and_field_completeness()
    test_period_infer_from_filenames()
    test_period_infer_from_date_columns()
    test_period_infer_from_to_banner_beats_row_dates()
    test_period_infer_filename_beats_cover_date_vote()
    test_quarter_months_order()
    test_class_label_mapping()
    test_banner_class_and_typos()
    test_ppn_share_bisola_style()
    test_currency_and_split_helpers()
    test_unitrust_foreign_title_stays_naira_with_currency_filter()
    test_adapter_registry_no_fallback()
    test_claims_leading_blank_default()
    test_claims_ppn_uses_percent_number_format()
    print("ALL UNIT TESTS PASSED")
