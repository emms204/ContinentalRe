"""IMPL-20260929-07: no false cedant_mismatch.

The selected cedant / broker is authoritative (a suspected mismatch is a
WARN only); detection reads only the file name, folder path and banner rows
above the header, and matches full registered names / codes as whole words.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cre_cleaner.adapters import ADAPTERS, _ALIASES
from cre_cleaner.core import batch as B
from cre_cleaner.core.batch import detect_cedants, detection_names, plan_batch, run_batch

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = REPO / "templates" / "TEMPLATE.xlsx"

# Insured names that contain other insurers' words (from the alias audit).
INSURED_TRAPS = [
    "ASABA MALL DEV/OLD MUTUAL INS PLC",
    "OLD MUTUAL NIGERIA GENERAL",
    "MUTUAL BENEFITS ASSURANCE STAFF",
    "REX PHARMACY LTD",
    "ROYAL EXCHANGE PLAZA TENANTS",
    "HEIRS HOLDINGS LIMITED",
    "UNITRUST STAFF CO-OP",
    "AXA MANSARD STAFF ESTATE",
    "CUSTODIAN AND ALLIED STAFF",
    "LASACO ASSURANCE HOUSE",
    "NEM HOUSE ANNEX",
    "CHI FARMS LIMITED",
]


def _listing(path: Path, *, banner=(), insureds=INSURED_TRAPS, header_row=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fire"
    for b in banner:
        ws.append([b] if isinstance(b, str) else list(b))
    if header_row:
        ws.append(["INSURED", "POLICY NO", "DEBIT NOTE", "PERIOD OF INSURANCE",
                   "SUM INSURED", "PREMIUM"])
    for i, name in enumerate(insureds):
        ws.append([name, f"P/{i}", f"DN{i}", "01/07/2020 - 30/06/2021", 1000, 10])
    wb.save(path)
    return path


# --- alias audit ------------------------------------------------------------

def test_single_word_aliases_are_disabled_for_detection():
    names, disabled = detection_names()
    dis = {a for a, _c, _r in disabled}
    assert "MUTUAL" in dis and "REX" in dis
    assert "MUTUAL" not in names and "REX" not in names
    for c, _b in ADAPTERS:                       # every registered name / code kept
        assert names[c] == c
    for a, c, _r in disabled:                    # only single-word fragments
        assert len(B._tok(a)) == 1 and a != c
    for (ac, _ab), (cc, _cb) in _ALIASES.items():  # the adapter lookup is untouched
        assert (ac, _ab) in _ALIASES


# --- C2: insured names never trigger a mismatch -----------------------------

@pytest.mark.parametrize("insured", INSURED_TRAPS)
def test_insured_names_do_not_trigger_mismatch(tmp_path, insured):
    f = _listing(tmp_path / "July Premium.xlsx", banner=[("OWN RETENTION", "TREATY")],
                 insureds=[insured])
    assert detect_cedants(f) == []
    plan = plan_batch([f], cedant="AIICO", broker="SCIB",
                      overrides={f.name: {"year": 2020, "quarter": 3}})
    assert not plan.pending and not plan.files[0].cedant_warning
    assert plan.groups[0].key == ("AIICO", "SCIB", 2020, 3)


def test_headerless_data_rows_are_not_scanned(tmp_path):
    # No recognisable header: the first tabular row ends the banner area.
    f = _listing(tmp_path / "x.xlsx", header_row=False)
    assert detect_cedants(f) == []


def test_old_mutual_in_a_banner_is_not_a_registered_cedant(tmp_path):
    f = _listing(tmp_path / "x.xlsx", banner=["OLD MUTUAL INSURANCE PLC"], insureds=["A"])
    assert detect_cedants(f) == []


# --- C3: a banner genuinely naming another cedant still warns ----------------

def test_banner_naming_other_cedant_warns_but_cleans_under_selection(tmp_path):
    f = _listing(tmp_path / "Q3 2020 premium.xlsx",
                 banner=["HEIRS INSURANCE LIMITED", "PREMIUM BORDEREAU JULY 2020"],
                 insureds=["ACME LTD"])
    assert detect_cedants(f) == ["HEIRS"]
    plan = plan_batch([f], cedant="AIICO", broker="SCIB")
    assert not plan.pending and plan.groups[0].key[:2] == ("AIICO", "SCIB")
    assert "HEIRS" in plan.files[0].cedant_warning


def test_selected_cedant_named_too_means_no_warning(tmp_path):
    f = _listing(tmp_path / "AIICO Q3 2020.xlsx",
                 banner=["AIICO INSURANCE PLC — reinsured with HEIRS"], insureds=["A"])
    plan = plan_batch([f], cedant="AIICO", broker="SCIB")
    assert not plan.files[0].cedant_warning


def test_folder_path_counts_as_evidence(tmp_path):
    d = tmp_path / "HEIRS" / "DIRECT" / "2020"
    d.mkdir(parents=True)
    f = _listing(d / "July Premium.xlsx", banner=[("OWN RETENTION", "TREATY")], insureds=["A"])
    assert detect_cedants(f) == ["HEIRS"]


def test_warn_reaches_the_group_exceptions_and_sidecar(tmp_path):
    from tests.test_impl_20260929_06 import _listing as prem_listing
    f = prem_listing(tmp_path / "2ND QTR 2024 HEIRS premium.xlsx", 2024, 2, n=2)
    plan = plan_batch([f], cedant="UNITRUST", broker="AGRIC")
    assert not plan.pending
    br = run_batch(plan, template=TEMPLATE, out_dir=tmp_path / "out",
                   work_dir=tmp_path / "work", base_dir=REPO)
    g = br.groups[0]
    assert g.status == "ok" and g.premium_rows == 2
    assert [e.reason for e in g.result.exceptions if e.severity == "WARN"
            and e.reason == "cedant_mismatch_suspected"] == ["cedant_mismatch_suspected"]
    assert g.warnings and not g.errors
    side = load_workbook(g.outputs[0]["exceptions_path"])
    text = " ".join(str(c) for ws in side.worksheets for r in ws.iter_rows(values_only=True)
                    for c in r if c)
    assert "cedant_mismatch_suspected" in text


# --- C4: no cedant selected → suggest, user confirms --------------------------

def test_no_selection_suggests_and_requires_confirmation(tmp_path):
    f = _listing(tmp_path / "HEIRS Q3 2020.xlsx", banner=["PREMIUM BORDEREAU"], insureds=["A"])
    plan = plan_batch([f], cedant=None, broker=None)
    fp = plan.files[0]
    assert fp.status == B.STATUS_CEDANT_UNCONFIRMED and fp.cedant == "HEIRS"
    assert fp.cedant_source == "suggested" and not plan.groups
    brk = next(b for c, b in ADAPTERS if c == "HEIRS")
    plan2 = plan_batch([f], cedant=None, broker=None,
                       overrides={f.name: {"cedant": "HEIRS", "broker": brk}})
    assert not plan2.pending and plan2.groups[0].key == ("HEIRS", brk, 2020, 3)


def test_no_selection_and_no_evidence_asks_user(tmp_path):
    f = _listing(tmp_path / "Q3 2020 premium.xlsx", banner=["PREMIUM BORDEREAU"], insureds=["A"])
    plan = plan_batch([f], cedant="", broker="")
    assert plan.files[0].status == B.STATUS_CEDANT_UNKNOWN and not plan.groups


def test_confirmed_cedant_contrary_to_evidence_still_warns(tmp_path):
    f = _listing(tmp_path / "HEIRS Q3 2020.xlsx", banner=["PREMIUM BORDEREAU"], insureds=["A"])
    plan = plan_batch([f], cedant=None, broker=None,
                      overrides={f.name: {"cedant": "AIICO", "broker": "SCIB"}})
    assert not plan.pending and plan.groups[0].key[:2] == ("AIICO", "SCIB")
    assert "HEIRS" in plan.files[0].cedant_warning
