"""RISK COV FROM / RISK COV TO map to Insurance Period From / To.

The 4-character partial-match guard stays: alias "TO" must not attach to
TOTAL, SUM INSURED TO DATE, AMOUNT DUE TO REINSURER, or a Unitrust share header.
"""
from __future__ import annotations

import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from cre_cleaner.core.map_columns import (
    CLAIMS_ALIASES,
    PREMIUM_ALIASES,
    map_simple_columns,
    missing_required_fields,
)
from cre_cleaner.core.quarterly import parse_premium_file
from cre_cleaner.core.reconcile import check_row_dates


def test_risk_cov_headers_map_for_premium_and_claims():
    headers = ["INSURED", "POLICY NO", "RISK COV FROM", "RISK COV TO", "PREMIUM"]
    for aliases in (PREMIUM_ALIASES, CLAIMS_ALIASES):
        cm = map_simple_columns(headers, aliases)
        assert cm.get("period_from") == 2
        assert cm.get("period_to") == 3


def test_short_to_does_not_match_unrelated_or_unitrust_headers():
    cm = map_simple_columns(
        ["INSURED", "TOTAL", "SUM INSURED TO DATE", "AMOUNT DUE TO REINSURER", "PREMIUM"],
        PREMIUM_ALIASES,
    )
    assert cm.get("period_to") is None
    assert cm.get("period_from") is None
    # "SUM INSURED" (11 chars) still prefix-matches this header; "TO" must not.
    assert cm.get("sum_insured") == 2
    assert "TOTAL" in cm.unmapped_headers
    assert "AMOUNT DUE TO REINSURER" in cm.unmapped_headers

    # Unitrust Agric: the share label is the band, and "TO" must not steal a column.
    unitrust = [
        "Policy No ", "Insured ", "Insurance Period ", "Sum Insured ", "Our share SI ",
        "Gross Premium ", "Our Share of G. Prem. ", "NET RETENTION ", "TREATY QUOTA SHARE",
    ]
    cm = map_simple_columns(unitrust, PREMIUM_ALIASES)
    assert cm.get("period") is not None
    assert cm.get("period_to") is None
    assert cm.get("period_from") is None

    exact = map_simple_columns(["INSURED", "TO", "FROM"], PREMIUM_ALIASES)
    assert exact.get("period_to") == 1
    assert exact.get("period_from") == 2


# September 2022 AIICO/ARK 'eng 2nd': COVER START / COVER END, not RISK COV.
# The 28 Sep alias list had neither name, so those two rows were left blank.
_SEPT_ENG_HEADER = [
    "INSURED", "POLICY NO", "DEBIT NOTE", "COVER START", "COVER END", "UW YEAR",
    "SUM INSURED", "PREMIUM", "RTNTN PPN", "RTNTN SUM INSURED", "RTNTN PREMIUM",
    "TREATY PPN1", "TREATY SUM INSURED", "TREATY PREMIUM1",
]


def test_sept_eng_cover_start_and_cover_end_map_to_period():
    cm = map_simple_columns(_SEPT_ENG_HEADER, PREMIUM_ALIASES)
    assert cm.get("period_from") == 3
    assert cm.get("period_to") == 4
    assert cm.get("uw_year") == 5

    wb = Workbook()
    ws = wb.active
    ws.title = "eng 2nd"
    ws.append(_SEPT_ENG_HEADER)
    ws.append([
        "JULIUS BERGER NIGERIA PLC", "207/00001/15/TQ/IB", "DRKD22003452",
        datetime(2022, 4, 1), datetime(2022, 12, 31), 2022, 8000000000, 14259435.3,
    ])
    ws.append([
        "NIGER DELTA POWER HOLDING CO. LTD.( ALAOJI PLANT)", "420/00002/15/TQ/HO",
        "CRHO22017264", datetime(2022, 7, 19), datetime(2022, 8, 19), 2022,
        -10995153750, -2012968.1,
    ])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Sept premium loc.xlsx"
        wb.save(path)
        rows, exc, _audit = parse_premium_file(path, "Sept")
    by_policy = {r.policy_no: r for r in rows}
    julius = by_policy["207/00001/15/TQ/IB"]
    niger = by_policy["420/00002/15/TQ/HO"]
    assert julius.period_from == datetime(2022, 4, 1)
    assert julius.period_to == datetime(2022, 12, 31)
    assert niger.period_from == datetime(2022, 7, 19)
    assert niger.period_to == datetime(2022, 8, 19)
    assert not any(e.reason in ("period_from_missing", "period_to_missing", "required_field_unmapped") for e in exc)


def test_unmapped_required_column_is_one_sheet_exception():
    wb = Workbook()
    ws = wb.active
    ws.title = "FIRE"
    ws.append(["INSURED", "POLICY NO", "PREMIUM", "MYSTERY COL"])
    ws.append(["Ada", "P1", 100, "not a date"])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "july premium.xlsx"
        wb.save(path)
        rows, exc, audit = parse_premium_file(path, "July")
    reasons = [e.reason for e in exc]
    assert reasons.count("required_field_unmapped") == 1
    gap = next(e for e in exc if e.reason == "required_field_unmapped")
    assert "period_from" in gap.detail and "period_to" in gap.detail
    assert gap.source_sheet == "FIRE"
    assert any(e.reason == "unmapped_headers" and "MYSTERY COL" in e.detail for e in exc)
    assert reasons.count("unmapped_headers") == 1
    assert "period_to_missing" not in reasons
    assert audit and "MYSTERY COL" in audit[0].notes
    cm = map_simple_columns(
        ["INSURED", "POLICY NO", "PREMIUM", "MYSTERY COL"], PREMIUM_ALIASES,
    )
    assert "MYSTERY COL" in cm.unmapped_headers
    assert "period_to" in missing_required_fields(cm, "premium")
    date_exc, _counts = check_row_dates(rows, [], [], year=2022, quarter=3)
    assert not any(e.reason == "period_to_missing" for e in date_exc)


def test_mapped_blank_to_still_warns_per_row_and_filled_to_does_not():
    wb = Workbook()
    ws = wb.active
    ws.title = "FIRE"
    ws.append(["INSURED", "POLICY NO", "PREMIUM", "RISK COV FROM", "RISK COV TO"])
    ws.append(["Ada", "207/00002/20/TQ/HO", 100, datetime(2022, 8, 25), datetime(2022, 12, 31)])
    ws.append(["Bo", "P2", 50, datetime(2022, 1, 1), None])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sept premium.xlsx"
        wb.save(path)
        rows, exc, _audit = parse_premium_file(path, "Sept")
    assert not any(e.reason == "required_field_unmapped" for e in exc)
    filled = next(r for r in rows if r.policy_no == "207/00002/20/TQ/HO")
    assert filled.period_from == datetime(2022, 8, 25)
    assert filled.period_to == datetime(2022, 12, 31)
    date_exc, _counts = check_row_dates(rows, [], [], year=2022, quarter=3)
    missing = [e for e in date_exc if e.reason == "period_to_missing"]
    assert len(missing) == 1
    assert "P2" in missing[0].detail
