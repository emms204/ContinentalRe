"""IMPL-20260929-06: multi-file batches — per-file period, one job per quarter,
isolation between jobs, pending (never dropped) files, delivery options."""
from __future__ import annotations

import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from cre_cleaner.adapters import ADAPTERS
from cre_cleaner.core import batch as B
from cre_cleaner.core.batch import plan_batch, run_batch
from cre_cleaner.io.delivery import (
    MERGED_BANNER, MERGED_SHEET, build_merged_workbook, data_sheet_counts, deliver,
)

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = REPO / "templates" / "TEMPLATE.xlsx"
CED, BRK = "UNITRUST", "AGRIC"
_MONTH = {1: "January", 2: "April", 3: "July", 4: "October"}
_END = {1: "March 31", 2: "June 30", 3: "September 30", 4: "December 31"}


def _banner_file(path: Path, year: int, q: int, n: int = 1, *, title="PREMIUM"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["TREATY RETURNS"])
    ws.append(["Period Date ( Effective Date )", f"From: [{_MONTH[q]} 01, {year} ]",
               f"To [ {_END[q]}, {year} ]"])
    ws.append(["S/NO", "INSURED", "START DATE", "SUM INSURED", "GROSS PREMIUM"])
    for i in range(n):
        ws.append([i + 1, f"INSURED {i}", datetime(year, 3 * q - 1, 5), 1000, 100 + i])
    wb.save(path)
    return path


# Premium listing in the layout the selected adapter parses (header row 1).
_HDR = ["Policy No ", "Insured ", "Policy Class ", "Risktype ", "Branch ", "Debit Note ",
        "Company's Share ", "MPL % ", "Insurance Period ", "Sum Insured ", "Our share SI ",
        "Gross Premium ", None, "Our Share of G. Prem. ", "NET RETENTION ", None, None,
        "TREATY QUOTA SHARE"]


def _listing(path: Path, year: int, q: int, n: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "AGRIC NAIRA"
    ws.append(_HDR)
    m = 3 * q - 1
    for i in range(n):
        ws.append([f"P{q}{i}/X", f"INSURED {i}", "Agriculture", "Poultry Insurance", "HEAD OFFICE",
                   f"DC{i}", 100, 100, f"{year}-{m:02d}-05 To {year + 1}-{m:02d}-04", 1000000,
                   1000000, 50000, None, 50000, 30, 300000, 15000, 70, 700000, 35000])
    wb.save(path)
    return path


def _plain_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["S/NO", "INSURED", "PREMIUM"])
    ws.append([1, "A", 5])
    wb.save(path)
    return path


# --- grouping ---------------------------------------------------------------

def test_files_keep_their_own_quarter_and_group(tmp_path):
    a = _banner_file(tmp_path / "premium one.xlsx", 2024, 2)
    b = _banner_file(tmp_path / "claims one.xlsx", 2024, 2, title="CLAIMS PAID")
    c = _banner_file(tmp_path / "other.xlsx", 2024, 4)
    plan = plan_batch([a, b, c], cedant=CED, broker=BRK)
    keys = [g.key for g in plan.groups]
    assert keys == [(CED, BRK, 2024, 2), (CED, BRK, 2024, 4)]
    assert sorted(f.name for f in plan.groups[0].files) == ["claims one.xlsx", "premium one.xlsx"]
    assert not plan.pending


def test_monthly_files_join_their_quarter(tmp_path):
    a = _banner_file(tmp_path / "APRIL PREMIUM.xlsx", 2023, 2)
    wb = Workbook()
    ws = wb.active
    ws.append(["S/NO", "INSURED", "PREMIUM"])
    ws.append([1, "A", 5])
    wb.save(tmp_path / "MAY PREMIUM 2023.xlsx")   # month + year in the name only
    plan = plan_batch([a, tmp_path / "MAY PREMIUM 2023.xlsx"], cedant=CED, broker=BRK)
    assert [g.key for g in plan.groups] == [(CED, BRK, 2023, 2)]
    assert len(plan.groups[0].files) == 2


def test_ambiguous_file_is_listed_not_dropped(tmp_path):
    a = _banner_file(tmp_path / "good.xlsx", 2024, 2)
    amb = _plain_file(tmp_path / "AGRIC BORDEREAUX.xlsx")
    plan = plan_batch([a, amb], cedant=CED, broker=BRK)
    assert [f.name for f in plan.pending] == ["AGRIC BORDEREAUX.xlsx"]
    assert plan.pending[0].status == B.STATUS_PERIOD_AMBIGUOUS
    assert {f.name for f in plan.files} == {"good.xlsx", "AGRIC BORDEREAUX.xlsx"}
    # the user sets the period → it joins that quarter's group
    plan2 = plan_batch([a, amb], cedant=CED, broker=BRK,
                       overrides={"AGRIC BORDEREAUX.xlsx": {"year": 2024, "quarter": 2}})
    assert not plan2.pending and len(plan2.groups[0].files) == 2
    assert {f.period_source for f in plan2.groups[0].files} == {"inferred", "override"}


def test_file_naming_another_cedant_warns_not_pending(tmp_path):
    # IMPL-20260929-07: the selected cedant is authoritative → WARN, not pending.
    other = next(c for c, _b in ADAPTERS if c != CED and " " not in c)
    f = _banner_file(tmp_path / f"{other} Q2 2024.xlsx", 2024, 2)
    plan = plan_batch([f], cedant=CED, broker=BRK)
    assert not plan.pending and plan.groups[0].key == (CED, BRK, 2024, 2)
    assert other in plan.files[0].detected_cedants and plan.files[0].cedant_warning
    ob = next(b for c, b in ADAPTERS if c == other)
    plan2 = plan_batch([f], cedant=CED, broker=BRK,
                       overrides={f.name: {"cedant": other, "broker": ob}})
    assert [g.key for g in plan2.groups] == [(other, ob, 2024, 2)]


def test_insured_old_mutual_does_not_cedant_mismatch(tmp_path):
    """AIICO (etc.) rows naming 'OLD MUTUAL INS PLC' must not trip MUTUAL alias."""
    path = tmp_path / "July Premium.xls"
    wb = Workbook()
    ws = wb.active
    ws.append(["OWN RETENTION", "TREATY", "INSURED", "POLICY NO", "PREMIUM"])
    ws.append([40000, 60000, "ASABA MALL DEV/OLD MUTUAL INS PLC", "P/1", 100])
    ws.append([10000, 20000, "MUTUAL BENEFITS ASSURANCE STAFF", "P/2", 50])
    wb.save(path)
    assert B.detect_cedants(path) == []
    plan = plan_batch([path], cedant="AIICO", broker="SCIB",
                      overrides={path.name: {"year": 2020, "quarter": 3}})
    assert not plan.pending
    assert plan.groups[0].key == ("AIICO", "SCIB", 2020, 3)


def test_filename_full_name_detects_but_single_word_alias_does_not(tmp_path):
    # IMPL-20260929-07: the single-word alias MUTUAL is disabled for detection;
    # the full registered name still detects — as a WARN, never pending.
    f = _banner_file(tmp_path / "MUTUAL Q2 2024 premium.xlsx", 2024, 2)
    assert B.detect_cedants(f) == []
    g = _banner_file(tmp_path / "MUTUAL BENEFITS Q2 2024 premium.xlsx", 2024, 2)
    assert "MUTUAL BENEFITS" in B.detect_cedants(g)
    plan = plan_batch([g], cedant="AIICO", broker="SCIB")
    assert not plan.pending and plan.files[0].cedant_warning


def test_sheet_title_mutual_benefits_still_detects(tmp_path):
    path = tmp_path / "July Premium.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["MUTUAL BENEFITS ASSURANCE PLC"])
    ws.append(["PREMIUM BORDEREAU Q3 2020"])
    ws.append(["INSURED", "POLICY NO", "GROSS PREMIUM"])
    ws.append(["Acme", "P1", 10])
    wb.save(path)
    assert B.detect_cedants(path) == ["MUTUAL BENEFITS"]


# --- isolation --------------------------------------------------------------

def test_failing_group_does_not_stop_others(tmp_path):
    a = _banner_file(tmp_path / "a.xlsx", 2024, 2)
    b = _banner_file(tmp_path / "b.xlsx", 2024, 3)
    plan = plan_batch([a, b], cedant=CED, broker=BRK)

    def runner(**kw):
        if kw["quarter"] == 2:
            raise RuntimeError("boom")
        return SimpleNamespace(exceptions=[], outputs=[{
            "output_path": str(tmp_path / "x.xlsx"), "premium_rows": 3,
            "claims_rows": 0, "outstanding_rows": 0}])
    br = run_batch(plan, template=TEMPLATE, out_dir=tmp_path / "out",
                   work_dir=tmp_path / "work", runner=runner)
    st = {g.group.quarter: g for g in br.groups}
    assert st[2].status == "failed" and "boom" in " ".join(st[2].errors)
    assert st[3].status == "ok" and st[3].premium_rows == 3
    assert [o["quarter"] for o in br.deliverable_outputs()] == [3]


def test_unreadable_file_group_reports_error_real_pipeline(tmp_path):
    good = _listing(tmp_path / "2ND QTR 2024 good.xlsx", 2024, 2, n=2)
    bad = tmp_path / "Q3 2024 PREMIUM.xlsx"
    bad.write_bytes(b"this is not a workbook")
    plan = plan_batch([good, bad], cedant=CED, broker=BRK)
    assert {g.key[2:] for g in plan.groups} == {(2024, 2), (2024, 3)}
    br = run_batch(plan, template=TEMPLATE, out_dir=tmp_path / "out",
                   work_dir=tmp_path / "work", base_dir=REPO)
    st = {g.group.quarter: g for g in br.groups}
    assert st[3].status == "failed" and st[3].errors
    assert st[2].status == "ok" and st[2].premium_rows == 2
    assert [Path(o["output_path"]).name for o in br.deliverable_outputs()] == [
        f"{CED}_{BRK}_2024_Q2_cleaned.xlsx"]


# --- delivery ---------------------------------------------------------------

def _cleaned(path: Path, title: str, q: int, n: int, *, extra_col=False):
    wb = Workbook()
    wb.active.title = "SUMMARY"
    ws = wb.create_sheet(title)
    ws["A1"] = f"CLASS — PREMIUM BORDEREAU — Q{q} 2024"
    hdr = ["S/NO", "POLICY NO.", "GROSS PREMIUM"] + (["NEW COL"] if extra_col else [])
    for j, h in enumerate(hdr, 1):
        ws.cell(3, j, h)
    for i in range(n):
        ws.append([i + 1, f"P{q}{i}", 10.0 * (i + 1)] + (["x"] if extra_col else []))
    wb.save(path)
    return {"output_path": str(path), "group": f"G Q{q}", "year": 2024, "quarter": q,
            "currency": "NGN", "source_files": [f"src{q}.xlsx"],
            "premium_rows": n, "claims_rows": 0, "outstanding_rows": 0}


def test_delivery_options(tmp_path):
    outs = [_cleaned(tmp_path / "A_Q2_cleaned.xlsx", "Class - PREMIUM", 2, 3),
            _cleaned(tmp_path / "A_Q4_cleaned.xlsx", "Class - PREMIUM", 4, 2, extra_col=True)]
    sep = deliver(outs, "separate", tmp_path / "d")
    assert [p.name for p in sep] == ["A_Q2_cleaned.xlsx", "A_Q4_cleaned.xlsx"]
    z = deliver(outs, "zip", tmp_path / "d", stem="T")[0]
    assert sorted(zipfile.ZipFile(z).namelist()) == ["A_Q2_cleaned.xlsx", "A_Q4_cleaned.xlsx"]
    m = deliver(outs, "merged", tmp_path / "d", stem="T")[0]
    assert "NOT_FOR_UPLOAD" in m.name
    assert data_sheet_counts(m) == {"Class - PREMIUM": 5}
    wb = load_workbook(m)
    assert wb.sheetnames[0] == MERGED_SHEET and wb[MERGED_SHEET]["A1"].value == MERGED_BANNER
    ws = wb["Class - PREMIUM"]
    assert "MERGED REVIEW" in ws["A1"].value and "2024 Q2 + 2024 Q4" in ws["A1"].value
    assert [ws.cell(r, 1).value for r in range(4, 9)] == [1, 2, 3, 4, 5]    # S/NO renumbered
    assert [ws.cell(r, 2).value for r in range(4, 9)] == ["P20", "P21", "P22", "P40", "P41"]
    rows = [[c.value for c in r] for r in wb[MERGED_SHEET].iter_rows(min_row=6)]
    assert [r[1] for r in rows] == ["2024 Q2", "2024 Q4"]            # source period marked
    assert rows[1][6] == "src4.xlsx" and rows[1][5] == "7–8"
    assert rows[1][8] == "NEW COL"                                    # reported, not hidden
    with pytest.raises(ValueError):
        deliver(outs, "bogus", tmp_path / "d")


def test_merged_keeps_currencies_apart(tmp_path):
    a = _cleaned(tmp_path / "A_cleaned.xlsx", "Class - PREMIUM", 2, 2)
    b = _cleaned(tmp_path / "A_FCY_cleaned.xlsx", "Class - PREMIUM", 2, 1)
    b["currency"] = "FCY"
    _p, rep = build_merged_workbook([a, b], tmp_path / "m.xlsx")
    assert rep["sheet_rows"] == {"Class - PREMIUM": 2, "Class - PREMIUM (FCY)": 1}


def test_batch_end_to_end_real_pipeline(tmp_path):
    files = [_listing(tmp_path / "2ND QTR 2024 a.xlsx", 2024, 2, n=3),
             _listing(tmp_path / "4TH QTR 2024 b.xlsx", 2024, 4, n=2)]
    br = run_batch(plan_batch(files, cedant=CED, broker=BRK), template=TEMPLATE,
                   out_dir=tmp_path / "out", work_dir=tmp_path / "work", base_dir=REPO)
    assert [(g.group.quarter, g.status, g.premium_rows) for g in br.groups] == [
        (2, "ok", 3), (4, "ok", 2)]
    for g in br.groups:
        assert not any(e.reason == "period_discovery_empty" for e in g.result.exceptions)
    outs = br.deliverable_outputs()
    per = {}
    for o in outs:
        for k, v in data_sheet_counts(Path(o["output_path"])).items():
            per[k] = per.get(k, 0) + v
    m = deliver(outs, "merged", tmp_path / "d")[0]
    assert data_sheet_counts(m) == per and sum(per.values()) == 5
