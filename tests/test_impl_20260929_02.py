"""IMPL-20260929-02: content-based table typing, secondary-source de-duplication,
claims period guard and the cross-quarter source check."""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest

from cre_cleaner.adapters import get_adapter
from cre_cleaner.adapters.base import claims_file_period_conflict
from cre_cleaner.core.table_type import (
    OTHER, OUTSTANDING, PAID, PREMIUM, UNKNOWN, classify_tab, table_header_rows,
)
from cre_cleaner.core.reconcile import overlap_counts
from cre_cleaner.models import ClaimsRow
from check_source_usage import shared_claims_sources

CLAIM_HDR = ["S/NO", "DATE OF LOSS", "NAME OF INSURED/CLAIMANT", "CLAIM NO", "POLICY NO",
             "COVER FROM", "COVER TO", "TOTAL CLAIMS PAID", "AIICO'S NET LIABILITY", "TREATY"]
CLAIM_HDR_PLAIN = [h.replace("TOTAL CLAIMS PAID", "TOTAL CLAIMS") for h in CLAIM_HDR]
OST_HDR = ["S/N", "LOSS DATE", "INSURED/CLAIMANT", "CLAIM NO", "POLICY NO", "PERIOD OF COVER",
           "U/YEAR", "SUM INSURED", "TOTAL OST RESERVE", "OWN SHARE.", "1ST SURPLUS "]
PREM_BAND = ["", "", "", "", "", "", "OWN RETENTION", "", "", "TREATY", "", ""]
PREM_HDR = ["S/N", "POLICY NO", "INSURED", "PERIOD FROM", "PERIOD TO", "SUM INSURED", "GROSS PREMIUM",
            "RTNTN PPN", "SUM INSURED1", "PREMIUM1", "TREATY PPN1", "TREATY SUM INSURED 1", "TREATY PREMIUM1"]


def _claim(i, amount=1000.0):
    return [i, datetime(2021, 11, 2), f"INSURED {i}", f"CL/0459{i:02d}/302/21/TQ/HO",
            f"302/000{i:02d}/18/TQ/HO", datetime(2021, 1, 1), datetime(2021, 12, 31),
            amount, amount * 0.1, amount * 0.9]


def _ost(i, amount=5000.0):
    return [i, "13/09/2021", f"INSURED {i}", f"CL/0263{i:02d}/503/17/TQ/HO", "503/00019/16/TQ/HO",
            "05/08/2021 to 04/11/2021", 2021, 1e6, amount, amount * 0.2, amount * 0.8]


def test_second_surplus_treaty_tab_two_paid_tables():
    # '2ND SURPLUS TREATY' has no type word; the sheet title covers both tables.
    rows = [["2ND SURP PAID CLAIM"], [], ["FIRE"], CLAIM_HDR, _claim(1), _claim(2), [2000.0], [],
            ["ENGINEERING"], CLAIM_HDR_PLAIN, _claim(3), [1000.0]]
    assert table_header_rows(rows) == [3, 9]
    tt = classify_tab("4TH Qtr 2021 Claims Bord.xls", "2ND SURPLUS TREATY", rows)
    assert tt.label == PAID and [t.label for t in tt.tables] == [PAID, PAID]


def test_claims_tab_in_premium_named_file_is_typed_by_content():
    rows = [["MISC. ACC PAID CLAIMS FOR OCTOBER 2021"], CLAIM_HDR, _claim(1), _claim(2)]
    tt = classify_tab("4th Qtr. 2021 - Premium ceded - LOCAL.xls", "MISC. ACC PAID OCT", rows)
    assert tt.label == PAID


@pytest.mark.parametrize("sheet", ["MARINE HULL OUSTANDING CLAIM", "FIRE OUTSANDING", "ENG OUTS",
                                   "MISC OST", "FIRE 2ND SURPLUS OUT. CLAIM"])
def test_outstanding_misspellings(sheet):
    rows = [["CLAIMS AS AT DECEMBER 2021"], [], OST_HDR, _ost(1), _ost(2)]
    assert classify_tab("x.xls", sheet, rows).label == OUTSTANDING


def test_reserve_header_is_outstanding_never_paid():
    rows = [[], OST_HDR, _ost(1), _ost(2)]
    tt = classify_tab("4th Qtr. 2020 - Claims Paid.xls", "MARINE HULL", rows)
    assert tt.label == OUTSTANDING


def test_paid_vs_outstanding_conflict_is_unknown():
    rows = [["OUTSTANDING CLAIMS AS AT 31 DEC 2021"], CLAIM_HDR, _claim(1), _claim(2)]
    tt = classify_tab("claims.xls", "FIRE PAID CLAIMS", rows)
    assert tt.label == UNKNOWN and any("conflict" in c for c in tt.conflicts)


def test_fac_oblig_premium_tab_inside_claims_file():
    hdr = ["S/N", "POLICY NO", "INSURED", "PERIOD OF INSURANCE", "SUM INSURED", "GROSS PREMIUM",
           "COMMISSION", "CEDED PREMIUM"]
    rows = [["FACULTATIVE OBLIGATORY PREMIUM UY 2019"], hdr,
            [1, "P/1", "A LTD", "2019-2020", 1e6, 5000, 500, 4500],
            [2, "P/2", "B LTD", "2019-2020", 2e6, 8000, 800, 7200]]
    assert classify_tab("1ST QTR 2020 - CLAIMS BORD.xlsx", "FAC OBLIG UY 2019", rows).label == PREMIUM


def test_rtntn_treaty_ppn_premium_tab():
    rows = [["FIRE PREMIUM JUNE 2023"], PREM_BAND, PREM_HDR,
            [1, "211/1", "X LTD", "1/6/23", "31/5/24", 1e8, 1e5, 25, 2.5e7, 2.5e4, 75, 7.5e7, 7.5e4]]
    assert classify_tab("JUNE loc.xlsx", "FIRE 1ST SURP", rows).label == PREMIUM


def test_both_families_is_unknown():
    hdr = ["POLICY NO", "INSURED", "CLAIM NO", "DATE OF LOSS", "GROSS PREMIUM", "SUM INSURED",
           "DEBIT NOTE", "COMMISSION"]
    rows = [hdr, ["P/1", "A", "CL/1", "1/1/21", 100, 1e6, "DN1", 10],
            ["P/2", "B", "CL/2", "1/1/21", 100, 1e6, "DN2", 10]]
    assert classify_tab("mixed.xlsx", "Sheet1", rows).label == UNKNOWN


def test_empty_statement_and_headerless_tabs_are_other():
    assert classify_tab("x.xls", "2nd surp cargo", [["2ND SURPLUS CARGO"], PREM_BAND, PREM_HDR]).label == OTHER
    assert classify_tab("x.xls", "STATEMENT OF ACCOUNT", [CLAIM_HDR, _claim(1)]).label == OTHER
    assert classify_tab("x.xls", "Sheet3", []).label == OTHER


def test_filename_is_only_a_tie_breaker():
    rows = [CLAIM_HDR_PLAIN, _claim(1), _claim(2)]
    assert classify_tab("3rd Qtr 2021 Outstanding.xls", "FIRE", rows).label == OUTSTANDING
    assert classify_tab("claims.xls", "FIRE", rows).label == UNKNOWN
    # an explicit tab/title/header word beats the file name
    rows_paid = [CLAIM_HDR, _claim(1), _claim(2)]
    assert classify_tab("3rd Qtr 2021 Outstanding.xls", "FIRE", rows_paid).label == PAID


@pytest.mark.parametrize("name,quarter,bad", [
    ("4TH Qtr 2021 Claims Bord.xls", 2, True),
    ("1ST QTR 2020 - CLAIMS BORD.xlsx", 2, True),
    ("2nd Qtr 2025 Claims Paid Bord.xlsx", 2, False),
    ("JANUARY CLAIMS 2021.xls", 2, True),
    ("MAY CLAIMS 2021.xls", 2, False),
    ("Claims Qtr 2021.xls", 2, False),   # no quarter in the name: not judged
])
def test_claims_file_period_guard(name, quarter, bad):
    assert (claims_file_period_conflict(name, quarter) is not None) is bad


def test_overlap_counts_multiset():
    a = ClaimsRow(claim_no="CL/1", policy_no="P", total_claims=100.0, amount_ret=10, amount_treaty=90)
    b = ClaimsRow(claim_no="CL/2", policy_no="P", total_claims=50.0)
    assert overlap_counts([a, b], [a, b]) == (2, 0)
    assert overlap_counts([a, a], [a]) == (1, 1)       # a repeated row is not a duplicate twice
    assert overlap_counts([b], []) == (0, 1)


def test_cross_quarter_row_explosion_is_caught():
    # every claims file of 2021 loaded into every 2021 quarter (paid 796 / OST 4,213 each)
    files = {"1st Qtr. 2021 - Claims Bord.xls": 1000, "2nd Qtr. 2021 - Claims Bord.xls": 1043,
             "3RD Qtr. 2021 - Claims Bord.xls": 957, "4TH Qtr 2021 Claims Bord.xls": 1011}
    usage = {f"AIICO_ARK_2021_Q{q}": dict(files) for q in range(1, 5)}
    assert set(shared_claims_sources(usage)) == set(files)
    ok = {f"AIICO_ARK_2021_Q{q}": {f: n} for q, (f, n) in enumerate(files.items(), 1)}
    assert shared_claims_sources(ok) == {}


def _write_claims_book(path: Path, tabs):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def test_duplicate_secondary_claims_file_is_not_double_loaded(tmp_path: Path):
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "2021"
    raw.mkdir()
    tabs = {"MISC. ACC PAID OCT": [["MISC ACCIDENT PAID CLAIMS OCTOBER 2021"], CLAIM_HDR,
                                   _claim(1, 1000.0), _claim(2, 2500.0)]}
    _write_claims_book(raw / "4TH Qtr 2021 Claims Bord.xlsx", tabs)
    _write_claims_book(raw / "4th Qtr. 2021 - Premium ceded - LOCAL.xlsx", tabs)
    tpl = ROOT / "templates" / "TEMPLATE.xlsx"
    if not tpl.exists():
        pytest.fail(f"template missing: {tpl}")
    res = run_pipeline(cedant="AIICO", broker="ARK", year=2021, quarter=4, raw_dir=raw,
                       template=tpl, out_dir=tmp_path / "out", bordereau_type="claims")
    assert len(res.claims_rows) == 2
    reasons = {e.reason for e in res.exceptions}
    assert "duplicate_source_not_loaded" in reasons
    sec = [a for a in res.source_audit if a.source_filename.startswith("4th Qtr. 2021 - Premium")]
    assert sec and all(a.sheet_type.startswith("secondary:") and a.rows_kept == 0 for a in sec)
    assert {a.detected_type for a in sec} == {"PAID"}
