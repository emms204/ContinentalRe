"""Cross-quarter source guard (IMPL-20260929-02).

Every claims / outstanding source file may feed exactly one quarter. A file
that feeds several quarters means a quarter loaded another quarter's claims
(the row explosion seen when every claims file of a year was loaded into every
quarter, e.g. paid 796 / outstanding 4,213 in each 2021 quarter).

Usage:  python tests/check_source_usage.py OUT_DIR [OUT_DIR ...]
Reads every *_source_audit.xlsx, exits 1 when a claims/outstanding source file
(rows kept > 0) is used by more than one quarter workbook.
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List

CLAIMS_SHEET_TYPES = {"paid", "outstanding"}
_STEM = re.compile(r"^(?P<label>.+?)_cleaned(?:_[A-Z]{3})?_source_audit$")


def shared_claims_sources(usage: Dict[str, Dict[str, int]]) -> Dict[str, List[str]]:
    """usage = {quarter_label: {source_filename: claims rows kept}} ->
    {source_filename: [quarter labels]} for files used by more than one quarter."""
    by_file: Dict[str, set] = defaultdict(set)
    for label, files in usage.items():
        for fname, kept in files.items():
            if kept:
                by_file[fname].add(label)
    return {f: sorted(q) for f, q in sorted(by_file.items()) if len(q) > 1}


def read_usage(out_dirs) -> Dict[str, Dict[str, int]]:
    from openpyxl import load_workbook
    usage: Dict[str, Dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for d in out_dirs:
        for p in sorted(Path(d).glob("*_source_audit.xlsx")):
            m = _STEM.match(p.stem)
            label = m.group("label") if m else p.stem
            wb = load_workbook(p, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            hdr = [str(h or "") for h in next(rows)]
            i_f, i_t, i_k = hdr.index("Source Filename"), hdr.index("Sheet Type"), hdr.index("Rows Kept")
            for r in rows:
                if r[i_t] in CLAIMS_SHEET_TYPES and (r[i_k] or 0) > 0:
                    usage[label][str(r[i_f])] += int(r[i_k])
            wb.close()
    return usage


def main(argv) -> int:
    if not argv:
        print(__doc__)
        return 2
    usage = read_usage(argv)
    shared = shared_claims_sources(usage)
    print(f"quarters={len(usage)} claims source files={len({f for u in usage.values() for f in u})}")
    for f, qs in shared.items():
        print(f"FAIL {f!r} feeds {len(qs)} quarters: {', '.join(qs)}")
    print("PASS: every claims/outstanding source file feeds one quarter" if not shared
          else f"FAIL: {len(shared)} claims source file(s) feed several quarters")
    return 1 if shared else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
