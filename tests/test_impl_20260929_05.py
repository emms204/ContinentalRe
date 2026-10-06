"""IMPL-20260929-05: period inference for single-file (demo upload) runs.

Covers: report-banner ranges (From…To / As At), whole-word header date
keywords, out-of-range date rejection, filename override (WARN
period_filename_override), no silent Q1 default (ERROR period_ambiguous),
single-file mode (no filename quarter gate, no discovery) and the
empty-discovery error for an inferred period.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from cre_cleaner.core import period_infer as P
from cre_cleaner.core.period_infer import (
    banner_range_period, header_has_phrase, infer_period, infer_period_from_labels,
)

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = REPO / "templates" / "TEMPLATE.xlsx"


def _save(path: Path, rows, title="PREMIUM"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


# --- banners ---------------------------------------------------------------

def test_banner_from_to_same_quarter():
    txt = ("\xa0Period Date \xa0( Effective Date ) \xa0From:\xa0[April 01, 2023 ]"
           " \xa0To\xa0[ June 30, 2023 ]")
    assert banner_range_period(txt) == (2023, 2, "range")
    assert banner_range_period("From 01/07/2024 To 30/09/2024") == (2024, 3, "range")


def test_banner_range_spanning_quarters_is_ignored():
    assert banner_range_period("From: [January 01, 2024 ] To [ December 31, 2024 ]") is None


def test_banner_as_at_single_date():
    assert banner_range_period("Period Date ( ) As At. [ September 30, 2024 ]") == (2024, 3, "as_at")


def test_print_date_row_is_not_a_banner():
    assert banner_range_period("Print Date [Monday 17th July 2023 10:15]") is None


def test_banner_beats_row_dates(tmp_path):
    rows = [
        ["TREATY RETURNS"],
        ["Period Date ( Effective Date )", "From: [April 01, 2023 ]", "To [ June 30, 2023 ]"],
        ["Print Date [Monday 17th July 2023]"],
        ["S/NO", "INSURED", "START DATE", "END DATE", "PREMIUM"],
        [1, "A", datetime(2023, 1, 5), datetime(2024, 1, 4), 100],
        [2, "B", datetime(2023, 2, 5), datetime(2024, 2, 4), 100],
        [3, "C", datetime(2023, 3, 5), datetime(2024, 3, 4), 100],
    ]
    _save(tmp_path / "mystery.xlsx", rows)
    r = infer_period(tmp_path)
    assert (r.year, r.quarter, r.confidence) == (2023, 2, "high")


# --- whole-word header keywords ---------------------------------------------

def test_header_keywords_match_whole_words():
    assert not header_has_phrase("TOTAL", "TO")
    assert not header_has_phrase("CLASS/SECTOR", "TO")
    assert header_has_phrase("COVER TO", "TO")
    assert header_has_phrase("Date Claim Paid", "DATE CLAIM PAID")
    assert P._header_date_weight("TOTAL") == 0.0
    assert P._header_date_weight("SECTOR") == 0.0


def test_dates_outside_1990_2100_are_ignored(tmp_path):
    rows = [["INSURED", "DATE OF LOSS", "TOTAL CLAIMS"]]
    rows += [["X", 87021 * 10, 1]] * 5  # serial far beyond 2100
    rows += [["A", datetime(2024, 5, 3), 1], ["B", datetime(2024, 6, 3), 1]]
    _save(tmp_path / "mystery.xlsx", rows, title="PAID")
    r = infer_period(tmp_path)
    assert r.year == 2024


# --- filename override / ambiguity -----------------------------------------

def test_filename_overrides_weak_content(tmp_path):
    # Only an old 'Date Raised'-style loss date in the content (Q2 2023) …
    rows = [["INSURED", "DATE OF LOSS", "OUTSTANDING"], ["A", datetime(2023, 4, 12), 10]]
    _save(tmp_path / "OUTSTANDING CLAIMS 2024 AS AT JUNE 2ND QTR.xlsx", rows, title="OUTSTANDING")
    r = infer_period(tmp_path)
    assert (r.year, r.quarter) == (2024, 2)
    assert r.filename_override and "file name" in r.filename_override


def test_high_confidence_content_is_not_overridden(tmp_path):
    rows = [["From: [April 01, 2023 ] To [ June 30, 2023 ]"],
            ["S/NO", "INSURED", "PREMIUM"], [1, "A", 1]]
    _save(tmp_path / "1ST QTR 2025.xlsx", rows)
    r = infer_period(tmp_path)
    assert (r.year, r.quarter) == (2023, 2) and not r.filename_override


def test_aiico_uy_banner_beats_stale_surplus_2015(tmp_path):
    """AIICO packs: SURPLUS still says '1st Quarter 2015'; STATEMENT UY is live."""
    from openpyxl import Workbook

    wb = Workbook()
    surplus = wb.active
    surplus.title = "SURPLUS"
    surplus.append(["1st Quarter 2015"])
    surplus.append(["INSURED", "DATE OF LOSS", "TOTAL CLAIMS"])
    surplus.append(["A", datetime(2015, 2, 1), 10])
    uy = wb.create_sheet("STATEMENT UY")
    uy.append(["Accounting Quarter: 3rd Quarter 2021"])
    uy.append(["INSURED", "DATE OF LOSS", "TOTAL CLAIMS"])
    uy.append(["B", datetime(2021, 8, 15), 20])
    path = tmp_path / "Paid Claims Bord Q3.xls.xlsx"
    wb.save(path)
    r = infer_period(path)
    assert (r.year, r.quarter) == (2021, 3)
    assert r.confidence == "high"


def test_filename_quarter_keeps_content_year_when_uy_missing(tmp_path):
    """Q4 packs often lack STATEMENT UY — filename Q + content year must win."""
    from openpyxl import Workbook

    wb = Workbook()
    surplus = wb.active
    surplus.title = "SURPLUS"
    surplus.append(["1st Quarter 2015"])
    surplus.append(["INSURED", "DATE OF LOSS", "TOTAL CLAIMS"])
    surplus.append(["A", datetime(2021, 11, 12), 10])
    surplus.append(["B", datetime(2021, 12, 3), 20])
    path = tmp_path / "Paid Claims Bord Q4..xlsx"
    wb.save(path)
    r = infer_period(path)
    assert (r.year, r.quarter) == (2021, 4)
    assert r.ok


def test_year_without_quarter_never_defaults_to_q1():
    r = infer_period_from_labels(["AGRIC BORDEREAUX 2024.xlsx"])
    assert r.year == 2024 and r.quarter is None and not r.ok


def test_ambiguous_period_is_an_error(tmp_path):
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "raw"
    raw.mkdir()
    _save(raw / "AGRIC BORDEREAUX.xlsx", [["S/NO", "INSURED", "PREMIUM"], [1, "A", 5]])
    res = run_pipeline(cedant="UNITRUST", broker="AGRIC", raw_dir=raw, template=TEMPLATE,
                       out_dir=tmp_path / "out", base_dir=REPO)
    reasons = {(e.severity, e.reason) for e in res.exceptions}
    assert ("ERROR", "period_ambiguous") in reasons
    assert not list((tmp_path / "out").glob("*_cleaned.xlsx"))


# --- single-file mode -------------------------------------------------------

def _q2_banner_file(path: Path):
    rows = [
        ["TREATY RETURNS"],
        ["Period Date ( Effective Date )", "From: [April 01, 2023 ]", "To [ June 30, 2023 ]"],
        ["S/NO", "INSURED", "START DATE", "SUM INSURED", "GROSS PREMIUM"],
        [1, "A", datetime(2023, 4, 5), 1000, 100],
    ]
    return _save(path, rows)


def test_single_file_mode_skips_filename_gate_and_discovery(tmp_path, monkeypatch):
    from cre_cleaner.adapters import get_adapter
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "raw"
    raw.mkdir()
    # File name says Q1; the banner says Q2 (high) → Q2 wins, file still processed.
    _q2_banner_file(raw / "1ST QTR CLAIMS 2023.xlsx")
    cls = type(get_adapter("UNITRUST", "AGRIC"))

    def boom(*a, **k):
        raise AssertionError("discovery must not run in single-file mode")
    monkeypatch.setattr(cls, "discover_claims_files", boom, raising=False)
    monkeypatch.setattr(cls, "discover_premium_files", boom, raising=False)
    res = run_pipeline(cedant="UNITRUST", broker="AGRIC", raw_dir=raw, template=TEMPLATE,
                       out_dir=tmp_path / "out", base_dir=REPO, single_file=True)
    reasons = {(e.severity, e.reason) for e in res.exceptions}
    assert (res.summary.get("year"), res.summary.get("quarter")) == (2023, 2)
    assert ("INFO", "single_file_mode") in reasons
    assert ("ERROR", "claims_file_out_of_period") not in reasons
    assert ("ERROR", "no_premium_files") not in reasons


def test_single_file_mode_is_automatic_for_one_inferred_file(tmp_path):
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "raw"
    raw.mkdir()
    _q2_banner_file(raw / "1ST QTR CLAIMS 2023.xlsx")
    res = run_pipeline(cedant="UNITRUST", broker="AGRIC", raw_dir=raw, template=TEMPLATE,
                       out_dir=tmp_path / "out", base_dir=REPO)
    assert any(e.reason == "single_file_mode" for e in res.exceptions)


def test_inferred_period_with_empty_discovery_is_an_error(tmp_path):
    from cre_cleaner.pipeline import run_pipeline
    raw = tmp_path / "raw"
    raw.mkdir()
    _q2_banner_file(raw / "a.xlsx")
    _q2_banner_file(raw / "b.xlsx")  # two files → folder mode, names match nothing
    res = run_pipeline(cedant="UNITRUST", broker="AGRIC", raw_dir=raw, template=TEMPLATE,
                       out_dir=tmp_path / "out", base_dir=REPO, single_file=False)
    reasons = {(e.severity, e.reason) for e in res.exceptions}
    assert ("ERROR", "period_discovery_empty") in reasons
    assert not list((tmp_path / "out").glob("*_cleaned.xlsx"))
